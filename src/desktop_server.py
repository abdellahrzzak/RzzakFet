import http.server
import json
import os
import io
import base64
import urllib.parse
import uuid
import threading
import time
import copy
from .models import InstitutionData, EducationalStructure, NonGeneralizedSubjectConfig, CustomRoom
from .curriculum_engine import CurriculumEngine
from .room_allocation_engine import RoomAllocationEngine
from .fet_xml_generator import FetXmlGenerator
from .timetable_engine import TimetableEngine
from .fet_native_bridge import FetNativeBridge
from .export_engine import ExportEngine
from .auth_engine import AuthEngine
from .exam_management_engine import ExamManagementEngine, DEFAULT_EXAM_PRESETS
from .golden_window_engine import GoldenWindowEngine
from .surveillance_daily_report_engine import SurveillanceDailyReportEngine
from .student_mobility_engine import mobility_engine
import sys
import webbrowser
from datetime import datetime

LATEST_PRINT_DOC = {"title": "وثيقة رسمية", "html": ""}

def get_base_dir():
    if getattr(sys, 'frozen', False):
        return sys._MEIPASS
    return os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

def get_data_dir():
    if getattr(sys, 'frozen', False):
        app_data = os.path.join(os.path.expanduser("~"), "AppData", "Local", "RzzakFet")
        os.makedirs(app_data, exist_ok=True)
        return os.path.join(app_data, "data")
    return os.path.join(get_base_dir(), "data")

BASE_DIR = get_base_dir()
DATA_DIR = get_data_dir()
UI_DIR = os.path.join(BASE_DIR, "ui")
ASSETS_DIR = os.path.join(BASE_DIR, "assets")
STATE_FILE = os.path.join(DATA_DIR, "app_state.json")

class RzzakFetAppState:
    def __init__(self):
        self.institution = InstitutionData()
        self.structure = EducationalStructure(classes_1_apic=12, classes_2_apic=10, classes_3_apic=10)
        self.curriculum_engine = CurriculumEngine(self.institution, self.structure)
        self.room_engine = RoomAllocationEngine(self.institution)
        self.fet_generator = FetXmlGenerator(self.institution, self.structure)
        self.timetable_engine = TimetableEngine()
        self.native_bridge = FetNativeBridge()
        self.auth_engine = AuthEngine()
        self.load_from_disk()

    def save_to_disk(self):
        try:
            os.makedirs(DATA_DIR, exist_ok=True)
            data = {
                "institution": {
                    "name": self.institution.institution_name,
                    "province": self.institution.province,
                    "academy": self.institution.academy,
                    "academic_year": self.institution.academic_year,
                    "gresa": self.institution.gresa_code,
                    "principal": self.institution.principal_name,
                    "general_rooms": self.institution.general_rooms_count,
                    "svt_labs": self.institution.svt_labs_count,
                    "pc_labs": self.institution.pc_labs_count,
                    "multimedia_rooms": self.institution.multimedia_rooms_count,
                    "sports_fields": self.institution.sports_fields_count,
                    "custom_room_types": self.institution.custom_room_types,
                    "break_time_slots": self.institution.break_time_slots,
                    "room_custom_names": self.institution.room_custom_names,
                    "teacher_custom_names": self.institution.teacher_custom_names,
                    "float_all_general_teachers": self.institution.float_all_general_teachers,
                    "subject_activity_constraints": self.institution.subject_activity_constraints,
                    "subject_split_modes": self.institution.subject_split_modes,
                    "teacher_counts_by_subject": self.institution.teacher_counts_by_subject,
                    "pe_wednesday_label": self.institution.pe_wednesday_label,
                    "general_wednesday_label": self.institution.general_wednesday_label,
                    "enable_wednesday_labels": self.institution.enable_wednesday_labels,
                    "morning_start_hour": self.institution.morning_start_hour,
                    "morning_period_times": self.institution.morning_period_times,
                    "afternoon_start_hour": self.institution.afternoon_start_hour,
                    "afternoon_period_times": self.institution.afternoon_period_times
                },
                "structure": {
                    "c1": self.structure.classes_1_apic,
                    "c2": self.structure.classes_2_apic,
                    "c3": self.structure.classes_3_apic,
                    "s1": self.structure.students_1_apic,
                    "s2": self.structure.students_2_apic,
                    "s3": self.structure.students_3_apic,
                },
                "non_generalized": [
                    {"name": cfg.name, "is_active": cfg.is_active}
                    for cfg in self.curriculum_engine.non_generalized_configs
                ],
                "manual_overrides": self.room_engine.manual_overrides,
                "constraints": self.fet_generator.all_available_constraints
            }
            with open(STATE_FILE, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
        except Exception as e:
            print("Error saving state to disk:", e)

    def load_from_disk(self):
        if not os.path.exists(STATE_FILE):
            return
        try:
            with open(STATE_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
            
            # Institution
            inst_data = data.get("institution", {})
            self.institution.institution_name = inst_data.get("name", self.institution.institution_name)
            self.institution.province = inst_data.get("province", self.institution.province)
            self.institution.academy = inst_data.get("academy", self.institution.academy)
            self.institution.academic_year = inst_data.get("academic_year", self.institution.academic_year)
            self.institution.gresa_code = inst_data.get("gresa", self.institution.gresa_code)
            self.institution.principal_name = inst_data.get("principal", self.institution.principal_name)
            self.institution.general_rooms_count = inst_data.get("general_rooms", self.institution.general_rooms_count)
            self.institution.svt_labs_count = inst_data.get("svt_labs", self.institution.svt_labs_count)
            self.institution.pc_labs_count = inst_data.get("pc_labs", self.institution.pc_labs_count)
            self.institution.multimedia_rooms_count = inst_data.get("multimedia_rooms", self.institution.multimedia_rooms_count)
            self.institution.sports_fields_count = inst_data.get("sports_fields", self.institution.sports_fields_count)
            self.institution.custom_room_types = inst_data.get("custom_room_types", [])
            self.institution.break_time_slots = inst_data.get("break_time_slots", self.institution.break_time_slots)
            self.institution.room_custom_names = inst_data.get("room_custom_names", {})
            self.institution.teacher_custom_names = inst_data.get("teacher_custom_names", {})
            self.institution.subject_activity_constraints = inst_data.get("subject_activity_constraints", {})
            if "subject_split_modes" in inst_data:
                self.institution.subject_split_modes.update(inst_data["subject_split_modes"])
            self.institution.teacher_counts_by_subject = inst_data.get("teacher_counts_by_subject", {})
            self.institution.pe_wednesday_label = inst_data.get("pe_wednesday_label", self.institution.pe_wednesday_label)
            self.institution.general_wednesday_label = inst_data.get("general_wednesday_label", self.institution.general_wednesday_label)
            self.institution.enable_wednesday_labels = inst_data.get("enable_wednesday_labels", self.institution.enable_wednesday_labels)
            self.institution.morning_start_hour = inst_data.get("morning_start_hour", self.institution.morning_start_hour)
            self.institution.morning_period_times = inst_data.get("morning_period_times", self.institution.morning_period_times)
            self.institution.afternoon_start_hour = inst_data.get("afternoon_start_hour", self.institution.afternoon_start_hour)
            self.institution.afternoon_period_times = inst_data.get("afternoon_period_times", self.institution.afternoon_period_times)

            # Structure
            struct_data = data.get("structure", {})
            self.structure.classes_1_apic = struct_data.get("c1", self.structure.classes_1_apic)
            self.structure.classes_2_apic = struct_data.get("c2", self.structure.classes_2_apic)
            self.structure.classes_3_apic = struct_data.get("c3", self.structure.classes_3_apic)
            self.structure.students_1_apic = struct_data.get("s1", self.structure.students_1_apic)
            self.structure.students_2_apic = struct_data.get("s2", self.structure.students_2_apic)
            self.structure.students_3_apic = struct_data.get("s3", self.structure.students_3_apic)

            # Non-generalized
            ng_data = {item["name"]: item["is_active"] for item in data.get("non_generalized", [])}
            for cfg in self.curriculum_engine.non_generalized_configs:
                if cfg.name in ng_data:
                    cfg.is_active = ng_data[cfg.name]

            # Room Overrides
            self.room_engine.manual_overrides = data.get("manual_overrides", {})

            # Constraints - preserve custom order, is_primary flag, activation, weight, param_val, and targeting
            saved_constraints = data.get("constraints", [])
            default_constraints = FetXmlGenerator._build_default_constraints()
            default_map = {c["id"]: c for c in default_constraints}
            default_primary_ids = {c["id"] for c in default_constraints if c.get("is_primary")}

            id_legacy_map = {
                "t_t_1": "tc_t_8",
                "t_t_2": "tc_t_10",
                "t_t_3": "tc_t_6",
                "t_t_4": "tc_t_4",
                "t_t_5": "tc_t_2",
                "t_t_6": "tc_t_12",
                "t_t_7": "tc_t_7",
                "t_t_8": "tc_t_1",
                "t_s_1": "tc_s_2",
                "t_s_2": "tc_s_3",
                "t_s_3": "tc_s_7",
                "t_s_4": "tc_s_9",
                "t_s_5": "tc_s_6",
                "t_s_6": "tc_s_5",
                "t_s_7": "tc_s_11",
                "t_s_8": "tc_s_1",
                "t_a_1": "tc_a_8",
                "t_all_1": "tc_base_2",
                "s_t_1": "sc_t_1",
                "s_sub_1": "sc_sub_2",
            }

            if saved_constraints:
                rebuilt_list = []
                seen_ids = set()

                for sc in saved_constraints:
                    raw_id = sc.get("id")
                    if not raw_id:
                        continue
                    c_id = id_legacy_map.get(raw_id, raw_id)
                    if c_id not in default_map or c_id in seen_ids:
                        continue
                    base = dict(default_map[c_id])
                    base["is_primary"] = sc.get("is_primary", c_id in default_primary_ids)
                    base["is_active"] = sc.get("is_active", base.get("is_active", True) if base["is_primary"] else False)
                    base["def_weight"] = float(sc.get("def_weight", base.get("def_weight", 100.0)))
                    if "param_val" in sc and base.get("param_name"):
                        base["param_val"] = sc["param_val"]
                    base["applies_to"] = sc.get("applies_to", base.get("applies_to", "all"))
                    base["selected_targets"] = sc.get("selected_targets", base.get("selected_targets", []))
                    rebuilt_list.append(base)
                    seen_ids.add(c_id)

                for def_c in default_constraints:
                    if def_c["id"] not in seen_ids:
                        rebuilt_list.append(dict(def_c))

                self.fet_generator.all_available_constraints = rebuilt_list
            else:
                self.fet_generator.all_available_constraints = [dict(c) for c in default_constraints]

        except Exception as e:
            print("Error loading state from disk:", e)

    def remap_timetable_teachers(self):
        cache_file = os.path.join(DATA_DIR, "timetable_result.json")
        if not os.path.exists(cache_file):
            return
        try:
            with open(cache_file, "r", encoding="utf-8") as f:
                tt = json.load(f)
            
            names_map = self.institution.teacher_custom_names
            if not names_map:
                return

            def map_tname(n):
                if not n:
                    return n
                return names_map.get(n, n)

            if "activities_dict" in tt:
                for aid, act in tt["activities_dict"].items():
                    act["teacher"] = map_tname(act.get("teacher", ""))

            if "all_activities_list" in tt:
                for act in tt["all_activities_list"]:
                    act["teacher"] = map_tname(act.get("teacher", ""))

            if "teacher_timetables" in tt:
                new_t_tt = {}
                for t, grid in tt["teacher_timetables"].items():
                    new_t_tt[map_tname(t)] = grid
                tt["teacher_timetables"] = new_t_tt

            if "student_timetables" in tt:
                for s_name, grid in tt["student_timetables"].items():
                    for d, hours in grid.items():
                        for h, cell in hours.items():
                            if cell and "teacher" in cell:
                                cell["teacher"] = map_tname(cell["teacher"])

            if "room_timetables" in tt:
                for r_name, grid in tt["room_timetables"].items():
                    for d, hours in grid.items():
                        for h, cell in hours.items():
                            if cell and "teacher" in cell:
                                cell["teacher"] = map_tname(cell["teacher"])

            if "master_teachers_data" in tt:
                for row in tt["master_teachers_data"]:
                    row["teacher"] = map_tname(row.get("teacher", ""))

            if "teachers_list" in tt:
                tt["teachers_list"] = sorted(list(set(map_tname(t) for t in tt["teachers_list"])))

            with open(cache_file, "w", encoding="utf-8") as f:
                json.dump(tt, f, ensure_ascii=False, indent=2)
        except Exception as e:
            print("Error remapping timetable teachers:", e)

    def get_full_state(self):
        quotas = self.curriculum_engine.get_subject_quotas()
        assignments = self.curriculum_engine.calculate_assignments()
        rooms = self.room_engine.allocate_rooms(assignments)
        all_teachers = [a.teacher_name for a in assignments if a.teacher_name]
        teachers_details = [{
            "name": a.teacher_name,
            "def_name": getattr(a, 'def_teacher_name', a.teacher_name),
            "real_name": self.institution.teacher_custom_names.get(getattr(a, 'def_teacher_name', a.teacher_name), ""),
            "subject": a.subject
        } for a in assignments if a.teacher_name]
        
        # Build default-to-custom teacher list for renaming UI and export
        all_default_teachers = []
        for a in assignments:
            if a.teacher_name:
                def_name = getattr(a, 'def_teacher_name', '') or a.teacher_name
                real_name = self.institution.teacher_custom_names.get(def_name, "")
                all_default_teachers.append({
                    "id": a.row_id,
                    "def_name": def_name,
                    "cur_name": a.teacher_name,
                    "real_name": real_name,
                    "subject": a.subject,
                    "total_hours": a.total_hours,
                    "assigned_classes": a.assigned_classes_str
                })
        
        all_classes_list = []
        for i in range(1, self.structure.classes_1_apic + 1):
            all_classes_list.append({"code": f"1APIC{i}", "label": f"1/{i}", "level": 1, "level_label": "الأولى إعدادي"})
        for i in range(1, self.structure.classes_2_apic + 1):
            all_classes_list.append({"code": f"2APIC{i}", "label": f"2/{i}", "level": 2, "level_label": "الثانية إعدادي"})
        for i in range(1, self.structure.classes_3_apic + 1):
            all_classes_list.append({"code": f"3APIC{i}", "label": f"3/{i}", "level": 3, "level_label": "الثالثة إعدادي"})
        
        # Build default-to-custom room list for renaming UI
        all_default_rooms = []
        for i in range(1, self.institution.general_rooms_count + 1):
            def_name = f"القاعة العامة رقم {i}"
            all_default_rooms.append({"def_name": def_name, "cur_name": self.institution.room_custom_names.get(def_name, def_name), "type": "قاعة عامة"})
        for j in range(self.institution.svt_labs_count):
            def_name = f"مختبر علوم الحياة والأرض {j+1}"
            all_default_rooms.append({"def_name": def_name, "cur_name": self.institution.room_custom_names.get(def_name, def_name), "type": "مختبر SVT"})
        for j in range(self.institution.pc_labs_count):
            def_name = f"مختبر الفيزياء والكيمياء {j+1}"
            all_default_rooms.append({"def_name": def_name, "cur_name": self.institution.room_custom_names.get(def_name, def_name), "type": "مختبر فيزياء"})
        for j in range(self.institution.multimedia_rooms_count):
            def_name = "قاعة الإعلاميات المتعددة الوسائط" if self.institution.multimedia_rooms_count == 1 else f"قاعة الإعلاميات {j+1}"
            all_default_rooms.append({"def_name": def_name, "cur_name": self.institution.room_custom_names.get(def_name, def_name), "type": "قاعة إعلاميات"})
        for j in range(self.institution.sports_fields_count):
            def_name = f"ملعب التربية البدنية {j+1}"
            all_default_rooms.append({"def_name": def_name, "cur_name": self.institution.room_custom_names.get(def_name, def_name), "type": "فضاء رياضي"})

        active_def_names = {t["def_name"] for t in all_default_teachers}
        active_teacher_custom_names = {
            k: v for k, v in self.institution.teacher_custom_names.items()
            if k in active_def_names
        }

        return {
            "institution": {
                "name": self.institution.institution_name,
                "province": self.institution.province,
                "academy": self.institution.academy,
                "academic_year": self.institution.academic_year,
                "gresa": self.institution.gresa_code,
                "principal": self.institution.principal_name,
                "general_rooms": self.institution.general_rooms_count,
                "svt_labs": self.institution.svt_labs_count,
                "pc_labs": self.institution.pc_labs_count,
                "multimedia_rooms": self.institution.multimedia_rooms_count,
                "sports_fields": self.institution.sports_fields_count,
                "total_rooms": self.institution.total_rooms,
                "max_room_capacity": self.institution.max_room_capacity_hours,
                "break_time_slots": self.institution.break_time_slots,
                "custom_room_types": self.institution.custom_room_types,
                "room_custom_names": self.institution.room_custom_names,
                "teacher_custom_names": active_teacher_custom_names,
                "float_all_general_teachers": self.institution.float_all_general_teachers,
                "teacher_counts_by_subject": self.institution.teacher_counts_by_subject,
                "morning_start_hour": self.institution.morning_start_hour,
                "morning_period_times": self.institution.morning_period_times,
                "afternoon_start_hour": self.institution.afternoon_start_hour,
                "afternoon_period_times": self.institution.afternoon_period_times,
                "all_default_rooms": all_default_rooms,
                "all_default_teachers": all_default_teachers
            },
            "structure": {
                "c1": self.structure.classes_1_apic,
                "c2": self.structure.classes_2_apic,
                "c3": self.structure.classes_3_apic,
                "s1": self.structure.students_1_apic,
                "s2": self.structure.students_2_apic,
                "s3": self.structure.students_3_apic,
                "total_classes": self.structure.total_classes,
                "total_students": self.structure.total_students,
                "utilization_pct": round((self.structure.total_classes / max(1, self.institution.total_rooms)) * 100, 1),
                "is_accommodated": self.structure.total_classes <= self.institution.total_rooms
            },
            "non_generalized": [
                {
                    "name": cfg.name,
                    "is_active": cfg.is_active,
                    "h1": cfg.hours_1_apic, "h2": cfg.hours_2_apic, "h3": cfg.hours_3_apic,
                    "room_type": cfg.room_type
                }
                for cfg in self.curriculum_engine.non_generalized_configs
            ],
            "quotas": [
                {
                    "name": q.name, "is_gen": q.is_generalized,
                    "h1": q.hours_1, "h2": q.hours_2, "h3": q.hours_3,
                    "total_hours": q.total_hours,
                    "req_rooms": q.required_rooms, "req_teachers": q.required_teachers,
                    "avail_teachers": self.institution.teacher_counts_by_subject.get(q.name, q.required_teachers),
                    "balance": self.institution.teacher_counts_by_subject.get(q.name, q.required_teachers) - q.required_teachers
                }
                for q in quotas
            ],
            "assignments": [
                {
                    "id": a.row_id, "subject": a.subject, "teacher": a.teacher_name,
                    "def_teacher": getattr(a, 'def_teacher_name', a.teacher_name),
                    "c1": a.classes_1, "c2": a.classes_2, "c3": a.classes_3,
                    "classes_str": a.assigned_classes_str,
                    "total_classes": a.total_classes, "levels_count": a.levels_count,
                    "two_levels_ok": a.two_levels_satisfied, "total_hours": a.total_hours,
                    "status": a.status_note
                }
                for a in assignments
            ],
            "rooms": [
                {
                    "id": r.room_id, "name": r.room_name, "type": r.room_type,
                    "subject": r.assigned_subject,
                    "m_teacher": r.morning_teacher, "e_teacher": r.afternoon_teacher,
                    "total_hours": r.total_hours, "max_cap": r.max_capacity_hours,
                    "util_pct": r.utilization_pct, "status": r.status,
                    "is_custom": r.is_custom
                }
                for r in rooms
            ],
            "all_constraints": self.fet_generator.all_available_constraints,
            "primary_constraints": self.fet_generator.get_primary_constraints(),
            "library_constraints": self.fet_generator.get_library_constraints(),
            "active_constraints": self.fet_generator.get_active_constraints(),
            "subject_activity_constraints": self.fet_generator.get_subject_activity_constraints(),
            "subject_split_modes": self.institution.subject_split_modes,
            "wednesday_settings": {
                "pe_label": self.institution.pe_wednesday_label,
                "general_label": self.institution.general_wednesday_label,
                "enabled": self.institution.enable_wednesday_labels
            },
            "teachers_list": all_teachers,
            "teachers_details": teachers_details,
            "classes_list": all_classes_list,
            "auth": self.auth_engine.get_auth_status(),
            "has_generated_timetable": os.path.exists(os.path.join(DATA_DIR, "timetable_result.json")),
            "fet_native": self.native_bridge.get_status(),
            "summary": {
                "total_teachers": len(assignments),
                "total_available_teachers": sum([self.institution.teacher_counts_by_subject.get(q.name, q.required_teachers) for q in quotas if q.total_hours > 0]),
                "total_classes": self.structure.total_classes,
                "total_rooms": self.institution.total_rooms,
                "total_hours": sum([a.total_hours for a in assignments]),
                "room_capacity": self.institution.max_room_capacity_hours
            }
        }

class ProgressiveSolverManager:
    def __init__(self, app_state):
        self.app_state = app_state
        self.is_running = False
        self.should_stop = False
        self.thread = None
        self.lock = threading.Lock()
        
        self.current_step = 0
        self.total_steps = 0
        self.current_constraint_name = ""
        self.stages = []
        self.best_timetable = None
        self.best_achieved_count = 0
        self.status_message = "جاهز للبدء"
        self.finished = False
        self.stopped_by_user = False
        self.start_time = 0
        self.elapsed_seconds = 0
        self.log_messages = []

    def get_status(self):
        with self.lock:
            elapsed = round(time.time() - self.start_time, 1) if self.is_running and self.start_time else self.elapsed_seconds
            pct = int((self.best_achieved_count / max(1, self.total_steps)) * 100) if self.total_steps > 0 else 0
            return {
                "is_running": self.is_running,
                "finished": self.finished,
                "stopped_by_user": self.stopped_by_user,
                "current_step": self.current_step,
                "total_steps": self.total_steps,
                "current_constraint_name": self.current_constraint_name,
                "best_achieved_count": self.best_achieved_count,
                "percentage": pct,
                "status_message": self.status_message,
                "elapsed_seconds": elapsed,
                "stages": self.stages,
                "has_timetable": self.best_timetable is not None,
                "log_messages": self.log_messages[-20:]
            }

    def start_generation(self):
        with self.lock:
            if self.is_running:
                return {"success": False, "error": "التوليد التراكمي قيد التشغيل بالفعل."}
            self.is_running = True
            self.should_stop = False
            self.finished = False
            self.stopped_by_user = False
            self.start_time = time.time()
            self.elapsed_seconds = 0
            self.best_timetable = None
            self.best_achieved_count = 0
            self.log_messages = ["🚀 بدء عملية إنتاج الجداول الدراسية قيداً بقيد..."]
            
            raw_constraints = self.app_state.fet_generator.all_available_constraints
            active_constraints = [c for c in raw_constraints if c.get("is_active", True)]

            active_codes = set(c.get("code") for c in active_constraints if c.get("code"))
            active_ids = set(c.get("id") for c in active_constraints if c.get("id"))
            active_all_identifiers = active_codes.union(active_ids)
            
            initial_stages = [
                {
                    "id": "base",
                    "name": "الجدول الهيكلي الأساسي (المواد، القاعات، عدم تجاوز 6 ساعات يومياً، وتفادي تكرار المادة)",
                    "codes": [
                        "ConstraintBreakTimes",
                        "ConstraintSubjectPreferredRooms",
                        "ConstraintMinDaysBetweenActivities",
                        "ConstraintStudentsMaxHoursDaily",
                        "ConstraintTeachersMaxHoursDaily"
                    ],
                    "status": "pending",
                    "desc": "تثبيت البنية الأساسية وإسنادات المواد، وتفادي تكرار المادة صباحاً ومساءً، وعدم تجاوز 6 ساعات يومياً",
                    "is_critical": True
                },
                {
                    "id": "max_continuous",
                    "name": "أقصى حصص متصلة (4 ساعات متصلة)",
                    "codes": ["ConstraintTeachersMaxHoursContinuously", "ConstraintTeachersMaxContinuousHours", "ConstraintStudentsMaxHoursContinuously", "ConstraintStudentsMaxContinuousHours"],
                    "status": "pending",
                    "desc": "منع إرهاق الأساتذة والتلاميذ بحصص متصلة مفرطة"
                },
                {
                    "id": "students_gaps",
                    "name": "منع الساعات الفارغة للمتعلمين (فجوات = 0)",
                    "codes": ["ConstraintStudentsMaxGapsPerDay", "ConstraintStudentsMaxGapsPerWeek", "ConstraintStudentsSetMaxGapsPerWeek"],
                    "status": "pending",
                    "desc": "جداول مستمرة للمتعلمين بدون أوقات فراغ وسط النهار"
                },
                {
                    "id": "teachers_gaps",
                    "name": "منع الساعات الفارغة للأساتذة (فجوات = 0)",
                    "codes": ["ConstraintTeachersMaxGapsPerDay", "ConstraintTeacherMaxGapsPerDay"],
                    "status": "pending",
                    "desc": "جداول مجمعة للأساتذة بدون ساعات انتظار فارغة"
                },
                {
                    "id": "teacher_home_rooms",
                    "name": "توطين الأساتذة في القاعات المخصصة (Home Rooms)",
                    "codes": ["ConstraintTeacherHomeRoom"],
                    "status": "pending",
                    "desc": "تثبيت الأساتذة في قاعاتهم الأم حسب طاقتها الاستيعابية"
                },
                {
                    "id": "max_mornings",
                    "name": "الحد الأقصى للفترات الصباحية للأساتذة (أقصى فترات صباحا)",
                    "codes": [
                        "ConstraintTeachersIntervalMaxDaysPerWeek",
                        "ConstraintTeacherIntervalMaxDaysPerWeek",
                        "tc_t_14",
                        "tc_t_13",
                        "ConstraintTeachersMaxMorningsPerWeek"
                    ],
                    "status": "pending",
                    "desc": "توزيع عادل للفترات الصباحية بين الأساتذة (4 فترات كحد أقصى)"
                }
            ]

            self.stages = []
            covered_codes = set()
            for st in initial_stages:
                if st.get("is_critical") or any(cd in active_all_identifiers for cd in st["codes"]):
                    self.stages.append(st)
                    for cd in st["codes"]:
                        covered_codes.add(cd)

            for c in active_constraints:
                code = c.get("code")
                cid = c.get("id")
                if (code and code not in covered_codes) and (not cid or cid not in covered_codes):
                    self.stages.append({
                        "id": cid or code,
                        "name": c.get("name", code),
                        "codes": [code] if code else [cid],
                        "status": "pending",
                        "desc": c.get("name", code)
                    })
                    if code:
                        covered_codes.add(code)
                    if cid:
                        covered_codes.add(cid)

            self.total_steps = len(self.stages)
            self.current_step = 0
            self.status_message = "جاري تحضير البنية الأساسية..."

            self.thread = threading.Thread(target=self._worker_loop, daemon=True)
            self.thread.start()
            return {"success": True, "message": "تم بدء التوليد التراكمي بنجاح."}

    def stop_generation(self):
        with self.lock:
            if not self.is_running:
                return {"success": True, "message": "المعالج متوقف حالياً."}
            self.should_stop = True
            self.stopped_by_user = True
            self.status_message = "تم الإيقاف بطلب من المستخدم واعتماد أفضل جدول محقق."
            self.log_messages.append("🛑 تم إيقاف البحث التراكمي بطلب من المستخدم.")

        if self.best_timetable:
            try:
                cache_file = os.path.join(DATA_DIR, "timetable_result.json")
                with open(cache_file, "w", encoding="utf-8") as f:
                    json.dump(self.best_timetable, f, ensure_ascii=False, indent=2)
            except Exception as e:
                print("Error saving stopped timetable:", e)
        return {"success": True, "message": "تم إيقاف التوليد واعتماد أفضل جدول محقق."}

    def _worker_loop(self):
        assignments = self.app_state.curriculum_engine.calculate_assignments()
        rooms = self.app_state.room_engine.allocate_rooms(assignments)
        
        cumulative_codes = set()
        c_by_code = {}
        for c in self.app_state.fet_generator.all_available_constraints:
            if c.get("code"):
                c_by_code[c["code"]] = dict(c)
            if c.get("id"):
                c_by_code[c["id"]] = dict(c)

        for idx, stage in enumerate(self.stages):
            if self.should_stop:
                break

            with self.lock:
                self.current_step = idx + 1
                self.current_constraint_name = stage["name"]
                stage["status"] = "running"
                self.status_message = f"جاري اختبار القيد ({idx + 1}/{self.total_steps}): {stage['name']}..."
                self.log_messages.append(f"⏳ مرحلة {idx + 1}/{self.total_steps}: اختبار {stage['name']}...")

            test_codes = set(cumulative_codes)
            for cd in stage["codes"]:
                test_codes.add(cd)

            active_list = []
            seen_ids = set()
            for cd in test_codes:
                if cd in c_by_code:
                    c_cand = c_by_code[cd]
                    cid = c_cand.get("id") or c_cand.get("code")
                    if cid not in seen_ids:
                        seen_ids.add(cid)
                        c_copy = dict(c_cand)
                        c_copy["is_active"] = True
                        active_list.append(c_copy)

            xml_data = self.app_state.fet_generator.generate_xml(assignments, rooms, active_constraints_override=active_list)

            t_start = time.time()
            res = self.app_state.timetable_engine.generate_timetable(xml_data, self.app_state.institution, rooms, timeout=60)
            dur = round(time.time() - t_start, 2)

            if res.get("success"):
                cumulative_codes = test_codes
                with self.lock:
                    stage["status"] = "success"
                    stage["duration"] = dur
                    self.best_timetable = res
                    self.best_achieved_count += 1
                    self.log_messages.append(f"✓ نجح تطبيق {stage['name']} في {dur}ث!")
                    try:
                        cache_file = os.path.join(DATA_DIR, "timetable_result.json")
                        with open(cache_file, "w", encoding="utf-8") as f:
                            json.dump(res, f, ensure_ascii=False, indent=2)
                    except Exception as e:
                        print("Error auto-saving progressive timetable:", e)
            else:
                with self.lock:
                    if stage.get("is_critical"):
                        stage["status"] = "failed"
                        self.log_messages.append(f"❌ تعذر حل البنية الأساسية ({dur}ث).")
                    else:
                        stage["status"] = "skipped"
                        stage["duration"] = dur
                        self.log_messages.append(f"⏭️ تم إرجاء {stage['name']} لصعوبته والاستمرار بأفضل جدول محقق ({dur}ث).")

        with self.lock:
            self.is_running = False
            self.finished = True
            self.elapsed_seconds = round(time.time() - self.start_time, 1)
            if self.stopped_by_user:
                self.status_message = f"تم الإيقاف! اعتُمد الجدول المحقق لـ ({self.best_achieved_count}/{self.total_steps}) قيود."
            elif self.best_achieved_count == self.total_steps:
                self.status_message = f"🎉 ممتاز! تم إنتاج الجدول النهائي محققاً جميع القيود بنسبة 100% في {self.elapsed_seconds}ث!"
            else:
                self.status_message = f"تم إنهاء التوليد التراكمي! اعتُمد أفضل جدول محقق لـ ({self.best_achieved_count}/{self.total_steps}) قيود في {self.elapsed_seconds}ث."
            self.log_messages.append("🏁 اكتملت معالجة التوليد التراكمي.")

def export_teachers_to_excel(app_state):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "هيئة التدريس"
        ws.sheet_view.rightToLeft = True

        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        
        real_col_fill = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
        real_col_font = Font(name="Segoe UI", size=11, bold=True, color="1E3A8A")
        
        data_font = Font(name="Segoe UI", size=10, bold=False, color="0F172A")
        data_font_bold = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
        
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )
        center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
        right_align = Alignment(horizontal='right', vertical='center', wrap_text=True)

        headers = [
            "الرقم",
            "الاسم الافتراضي",
            "الاسم الحقيقي (اكتب أو انسخ هنا)",
            "المادة",
            "الأقسام المسندة",
            "مجموع الحصص"
        ]
        
        ws.row_dimensions[1].height = 32
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border

        assignments = app_state.curriculum_engine.calculate_assignments()
        row_num = 2
        for a in assignments:
            if not a.teacher_name:
                continue
            def_name = getattr(a, 'def_teacher_name', '') or a.teacher_name
            real_name = app_state.institution.teacher_custom_names.get(def_name, "")
            
            ws.row_dimensions[row_num].height = 24

            c1 = ws.cell(row=row_num, column=1, value=a.row_id)
            c1.alignment = center_align
            c1.font = data_font
            c1.border = thin_border

            c2 = ws.cell(row=row_num, column=2, value=def_name)
            c2.alignment = right_align
            c2.font = data_font_bold
            c2.border = thin_border

            c3 = ws.cell(row=row_num, column=3, value=real_name)
            c3.alignment = right_align
            c3.font = real_col_font
            c3.fill = real_col_fill
            c3.border = thin_border

            c4 = ws.cell(row=row_num, column=4, value=a.subject)
            c4.alignment = center_align
            c4.font = data_font
            c4.border = thin_border

            c5 = ws.cell(row=row_num, column=5, value=a.assigned_classes_str)
            c5.alignment = center_align
            c5.font = data_font
            c5.border = thin_border

            c6 = ws.cell(row=row_num, column=6, value=f"{a.total_hours}س")
            c6.alignment = center_align
            c6.font = data_font_bold
            c6.border = thin_border

            row_num += 1

        col_widths = {1: 8, 2: 28, 3: 35, 4: 24, 5: 35, 6: 16}
        for col_idx, width in col_widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        buf = io.BytesIO()
        wb.save(buf)
        file_bytes = buf.getvalue()
    except Exception as e:
        print("Excel generation fallback to CSV due to error:", e)
        import csv
        buf = io.StringIO()
        writer = csv.writer(buf, delimiter=';')
        writer.writerow(["الرقم", "الاسم الافتراضي", "الاسم الحقيقي (اكتب أو انسخ هنا)", "المادة", "الأقسام المسندة", "مجموع الحصص"])
        assignments = app_state.curriculum_engine.calculate_assignments()
        for a in assignments:
            if not a.teacher_name: continue
            def_name = getattr(a, 'def_teacher_name', '') or a.teacher_name
            real_name = app_state.institution.teacher_custom_names.get(def_name, "")
            writer.writerow([a.row_id, def_name, real_name, a.subject, a.assigned_classes_str, f"{a.total_hours}س"])
        file_bytes = ('\ufeff' + buf.getvalue()).encode('utf-8')

    desktop_dir = os.path.join(os.path.expanduser("~"), "Desktop", "RzzakFet")
    os.makedirs(desktop_dir, exist_ok=True)
    dest_path = os.path.join(desktop_dir, "لائحة_أسماء_الأساتذة_RzzakFet.xlsx")
    try:
        with open(dest_path, "wb") as f:
            f.write(file_bytes)
    except Exception as ex:
        print("Could not write Excel to Desktop:", ex)

    return file_bytes, dest_path

def import_teachers_from_file(app_state, file_bytes: bytes, filename: str) -> dict:
    if not file_bytes:
        return {"success": False, "error": "لم يتم استلام أي بيانات في الملف."}

    imported = {}
    assignments = app_state.curriculum_engine.calculate_assignments()
    known_defs = [getattr(a, 'def_teacher_name', a.teacher_name) for a in assignments if a.teacher_name]

    is_excel = filename.lower().endswith(('.xlsx', '.xlsm', '.xltx')) or (len(file_bytes) > 4 and file_bytes[:4] == b'PK\x03\x04')
    rows = []

    if is_excel:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
        except Exception as e:
            print("Error parsing Excel with openpyxl:", e)

    if not rows:
        for enc in ['utf-8-sig', 'utf-8', 'cp1256', 'latin-1']:
            try:
                text = file_bytes.decode(enc)
                break
            except Exception:
                text = None
        if text is None:
            text = file_bytes.decode('utf-8', errors='ignore')
        
        import csv
        delimiter = ';' if ';' in text else (',' if ',' in text else '\t')
        rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))

    if not rows:
        return {"success": False, "error": "تعذر قراءة محتوى الملف أو الملف فارغ."}

    def_col_idx = 1
    real_col_idx = 2

    header_row = [str(c or "").strip() for c in rows[0]]
    for idx, h in enumerate(header_row):
        if "افتراض" in h:
            def_col_idx = idx
        elif "حقيق" in h:
            real_col_idx = idx

    if len(header_row) == 2:
        def_col_idx = 0
        real_col_idx = 1

    start_row = 1 if any("اسم" in h or "مادة" in h or "افتراض" in h or "حقيق" in h for h in header_row) else 0

    row_cursor = 0
    for r_i in range(start_row, len(rows)):
        row = rows[r_i]
        if not row or not any(row):
            continue

        col_def = str(row[def_col_idx]).strip() if len(row) > def_col_idx and row[def_col_idx] is not None else ""
        col_real = str(row[real_col_idx]).strip() if len(row) > real_col_idx and row[real_col_idx] is not None else ""

        if not col_real and len(row) >= 2 and col_def:
            alt_real = str(row[1]).strip() if row[1] is not None else ""
            if alt_real and alt_real != col_def:
                col_real = alt_real

        if not col_real or col_real in ["الاسم الحقيقي", "الاسم الحقيقي (اكتب أو انسخ هنا)", "-", "None"]:
            row_cursor += 1
            continue

        matched_def = None
        if col_def in known_defs:
            matched_def = col_def
        else:
            for kd in known_defs:
                if col_def and (col_def in kd or kd in col_def):
                    matched_def = kd
                    break

        if not matched_def and row_cursor < len(known_defs):
            matched_def = known_defs[row_cursor]

        if matched_def and col_real:
            imported[matched_def] = col_real

        row_cursor += 1

    if imported:
        app_state.institution.teacher_custom_names.update(imported)
        app_state.save_to_disk()
        app_state.curriculum_engine.calculate_assignments()
        app_state.remap_timetable_teachers()

    return {
        "success": True,
        "count": len(imported),
        "names": imported
    }

APP_STATE = RzzakFetAppState()
PROGRESSIVE_MANAGER = ProgressiveSolverManager(APP_STATE)
EXPORT_ENGINE = ExportEngine()

class RzzakFetHandler(http.server.SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=UI_DIR, **kwargs)

    def send_json_data(self, data_obj, status=200):
        raw = json.dumps(data_obj, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(raw)))
        self.send_header("Cache-Control", "no-cache, no-store, must-revalidate, max-age=0")
        self.send_header("Pragma", "no-cache")
        self.send_header("Expires", "0")
        self.send_header("Access-Control-Allow-Origin", "*")
        self.end_headers()
        self.wfile.write(raw)
        try:
            self.wfile.flush()
        except Exception:
            pass

    def send_html_data(self, html_text, status=200):
        raw = html_text.encode("utf-8") if isinstance(html_text, str) else html_text
        self.send_response(status)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(raw)))
        self.send_header("Cache-Control", "no-cache, no-store, must-revalidate, max-age=0")
        self.send_header("Pragma", "no-cache")
        self.send_header("Expires", "0")
        self.send_header("Access-Control-Allow-Origin", "*")
        self.end_headers()
        self.wfile.write(raw)
        try:
            self.wfile.flush()
        except Exception:
            pass

    def end_headers(self):
        self.send_header("Cache-Control", "no-cache, no-store, must-revalidate, max-age=0")
        self.send_header("Pragma", "no-cache")
        self.send_header("Expires", "0")
        self.send_header("Access-Control-Allow-Origin", "*")
        super().end_headers()

    def do_GET(self):
        parsed = urllib.parse.urlparse(self.path)
        query_params = urllib.parse.parse_qs(parsed.query)

        if parsed.path == "/print_preview":
            html_content = LATEST_PRINT_DOC.get("html", "<h1 style='text-align:center;padding:50px;'>لا توجد وثيقة حالية للمعاينة</h1>")
            doc_title = LATEST_PRINT_DOC.get("title", "معاينة الوثيقة الرسمية")
            if "<body" in html_content and "window.print()" not in html_content:
                html_content = html_content.replace("<body", '<body onload="setTimeout(function(){window.print();}, 600)"')
            encoded = html_content.encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Cache-Control", "no-cache, no-store, must-revalidate")
            self.send_header("Content-Length", str(len(encoded)))
            self.end_headers()
            self.wfile.write(encoded)
            return

        if parsed.path in ("/", "/index.html"):
            candidates = [
                os.path.join(DATA_DIR, "ui", "index.html"),
                os.path.join(UI_DIR, "index.html"),
                os.path.join(os.getcwd(), "ui", "index.html"),
                os.path.join(os.environ.get("LOCALAPPDATA", ""), "Programs", "RzzakFet", "ui", "index.html"),
                os.path.join(os.environ.get("LOCALAPPDATA", ""), "Programs", "RzzakFet", "_internal", "ui", "index.html"),
            ]
            valid_candidates = [p for p in candidates if os.path.exists(p) and os.path.getsize(p) > 10000]
            if valid_candidates:
                index_path = max(valid_candidates, key=os.path.getmtime)
            else:
                index_path = os.path.join(UI_DIR, "index.html")
            if os.path.exists(index_path):
                with open(index_path, "rb") as f:
                    content = f.read()
                self.send_response(200)
                self.send_header("Content-Type", "text/html; charset=utf-8")
                self.send_header("Cache-Control", "no-cache, no-store, must-revalidate, max-age=0")
                self.send_header("Pragma", "no-cache")
                self.send_header("Expires", "0")
                self.send_header("Content-Length", str(len(content)))
                self.end_headers()
                self.wfile.write(content)
                return

        if parsed.path.startswith("/assets/"):
            asset_subpath = parsed.path[len("/assets/"):].lstrip("/")
            asset_full = os.path.join(ASSETS_DIR, asset_subpath)
            if os.path.exists(asset_full) and os.path.isfile(asset_full):
                self.send_response(200)
                if asset_full.endswith(".ico"):
                    self.send_header("Content-Type", "image/x-icon")
                elif asset_full.endswith(".png"):
                    self.send_header("Content-Type", "image/png")
                elif asset_full.endswith(".svg"):
                    self.send_header("Content-Type", "image/svg+xml")
                elif asset_full.endswith(".css"):
                    self.send_header("Content-Type", "text/css")
                elif asset_full.endswith(".js"):
                    self.send_header("Content-Type", "application/javascript")
                else:
                    self.send_header("Content-Type", "application/octet-stream")
                with open(asset_full, "rb") as f:
                    file_content = f.read()
                self.send_header("Content-Length", str(len(file_content)))
                self.end_headers()
                self.wfile.write(file_content)
                return
            else:
                self.send_error(404, "Asset not found")
                return
        elif parsed.path == "/favicon.ico":
            ico_path = os.path.join(ASSETS_DIR, "icon.ico")
            if os.path.exists(ico_path):
                with open(ico_path, "rb") as f:
                    ico_content = f.read()
                self.send_response(200)
                self.send_header("Content-Type", "image/x-icon")
                self.send_header("Content-Length", str(len(ico_content)))
                self.end_headers()
                self.wfile.write(ico_content)
                return

        if parsed.path == "/api/ping":
            self.send_json_data({"status": "ok", "app": "RzzakFet", "ready": True})
            return
        elif parsed.path == "/api/state":
            self.send_json_data(APP_STATE.get_full_state())
        elif parsed.path == "/api/auth/status":
            force = query_params.get("force", ["0"])[0] in ("1", "true")
            self.send_json_data(APP_STATE.auth_engine.get_auth_status(force_cloud=force))
        elif parsed.path == "/api/app/check_updates":
            force = query_params.get("force", ["0"])[0] in ("1", "true")
            self.send_json_data(APP_STATE.auth_engine.check_for_updates(force_cloud=force))
        elif parsed.path == "/api/admin/list_users":
            auth_info = APP_STATE.auth_engine.get_auth_status()
            if auth_info.get("is_admin"):
                self.send_json_data({"success": True, "users": APP_STATE.auth_engine.list_all_users()})
            else:
                self.send_json_data({"success": False, "error": "غير مصرح لك بالوصول لهذه البيانات."})
        elif parsed.path == "/api/progressive_status":
            self.send_json_data(PROGRESSIVE_MANAGER.get_status())
        elif parsed.path == "/api/export_teachers_excel":
            file_bytes, dest_path = export_teachers_to_excel(APP_STATE)
            self.send_response(200)
            self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.send_header("Content-Disposition", 'attachment; filename="teachers_names_RzzakFet.xlsx"')
            self.send_header("Content-Length", str(len(file_bytes)))
            self.end_headers()
            self.wfile.write(file_bytes)
            return
        elif parsed.path == "/api/export_fet":
            self.send_response(200)
            self.send_header("Content-Type", "application/xml; charset=utf-8")
            self.send_header("Content-Disposition", f"attachment; filename=RzzakFet_{APP_STATE.institution.gresa_code}.fet")
            self.end_headers()
            assignments = APP_STATE.curriculum_engine.calculate_assignments()
            rooms = APP_STATE.room_engine.allocate_rooms(assignments)
            xml_data = APP_STATE.fet_generator.generate_xml(assignments, rooms)
            self.wfile.write(xml_data.encode("utf-8"))
        elif parsed.path == "/api/get_timetable":
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            tt_data = APP_STATE.timetable_engine.get_cached_timetable() or {"success": False, "error": "لم يتم إنتاج أي جدول بعد."}
            self.wfile.write(json.dumps(tt_data, ensure_ascii=False).encode("utf-8"))
        elif parsed.path == "/api/fet_status":
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            self.wfile.write(json.dumps(APP_STATE.native_bridge.get_status(), ensure_ascii=False).encode("utf-8"))
        elif parsed.path == "/api/open_fet_folder":
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = APP_STATE.native_bridge.open_working_directory()
            self.wfile.write(json.dumps(res, ensure_ascii=False).encode("utf-8"))
        elif parsed.path == "/api/open_desktop_folder":
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = EXPORT_ENGINE.open_desktop_folder()
            self.wfile.write(json.dumps(res, ensure_ascii=False).encode("utf-8"))
        elif parsed.path == "/api/exams/presets":
            self.send_json_data({"success": True, "presets": DEFAULT_EXAM_PRESETS})
        elif parsed.path == "/api/golden_window/analyze":
            tt_data = APP_STATE.timetable_engine.get_cached_timetable() or {}
            assignments = APP_STATE.curriculum_engine.calculate_assignments()
            engine = GoldenWindowEngine(tt_data)
            res = engine.analyze_all_subjects(assignments)
            self.send_json_data(res)
        elif parsed.path == "/api/surveillance/day_summary":
            day_name = query_params.get("day", ["الاثنين"])[0]
            tt_data = APP_STATE.timetable_engine.get_cached_timetable() or {}
            engine = SurveillanceDailyReportEngine(tt_data, APP_STATE.institution)
            res = engine.get_scheduled_day_summary(day_name)
            self.send_json_data({"success": True, "data": res})
        elif parsed.path == "/api/mobility/summary":
            try:
                mobility_engine.sync_with_institution_and_structure(APP_STATE.institution, APP_STATE.structure)
            except Exception as e:
                print(f"[Mobility Sync Error]: {e}")
            self.send_json_data(mobility_engine.compute_mobility_analysis())
        elif parsed.path == "/api/mobility/export_excel":
            excel_io = mobility_engine.generate_excel_export()
            content = excel_io.getvalue()
            self.send_response(200)
            self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.send_header("Content-Disposition", 'attachment; filename="mobility_report.xlsx"')
            self.send_header("Content-Length", str(len(content)))
            self.end_headers()
            self.wfile.write(content)
        elif parsed.path == "/export/all_classes":
            tt_data = APP_STATE.timetable_engine.get_cached_timetable()
            if not tt_data:
                assignments = APP_STATE.curriculum_engine.calculate_assignments()
                rooms = APP_STATE.room_engine.allocate_rooms(assignments)
                xml_data = APP_STATE.fet_generator.generate_xml(assignments, rooms)
                tt_data = APP_STATE.timetable_engine.generate_timetable(xml_data, APP_STATE.institution, rooms)
            
            exp_res = EXPORT_ENGINE.export_all(tt_data, APP_STATE.get_full_state())
            with open(exp_res["all_classes_file"], "r", encoding="utf-8") as f:
                content = f.read()
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.end_headers()
            self.wfile.write(content.encode("utf-8"))
        elif parsed.path == "/export/all_teachers":
            tt_data = APP_STATE.timetable_engine.get_cached_timetable()
            if not tt_data:
                assignments = APP_STATE.curriculum_engine.calculate_assignments()
                rooms = APP_STATE.room_engine.allocate_rooms(assignments)
                xml_data = APP_STATE.fet_generator.generate_xml(assignments, rooms)
                tt_data = APP_STATE.timetable_engine.generate_timetable(xml_data, APP_STATE.institution, rooms)
            
            exp_res = EXPORT_ENGINE.export_all(tt_data, APP_STATE.get_full_state())
            with open(exp_res["all_teachers_file"], "r", encoding="utf-8") as f:
                content = f.read()
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.end_headers()
            self.wfile.write(content.encode("utf-8"))
        elif parsed.path == "/export/master":
            tt_data = APP_STATE.timetable_engine.get_cached_timetable()
            if not tt_data:
                assignments = APP_STATE.curriculum_engine.calculate_assignments()
                rooms = APP_STATE.room_engine.allocate_rooms(assignments)
                xml_data = APP_STATE.fet_generator.generate_xml(assignments, rooms)
                tt_data = APP_STATE.timetable_engine.generate_timetable(xml_data, APP_STATE.institution, rooms)
            
            exp_res = EXPORT_ENGINE.export_all(tt_data, APP_STATE.get_full_state())
            with open(exp_res["master_file"], "r", encoding="utf-8") as f:
                content = f.read()
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.end_headers()
            self.wfile.write(content.encode("utf-8"))
        elif parsed.path == "/export/single":
            tt_data = APP_STATE.timetable_engine.get_cached_timetable()
            if not tt_data:
                assignments = APP_STATE.curriculum_engine.calculate_assignments()
                rooms = APP_STATE.room_engine.allocate_rooms(assignments)
                xml_data = APP_STATE.fet_generator.generate_xml(assignments, rooms)
                tt_data = APP_STATE.timetable_engine.generate_timetable(xml_data, APP_STATE.institution, rooms)
            
            cat = query_params.get("type", ["class"])[0]
            name = query_params.get("name", [""])[0]

            grid = {}
            if cat == "teacher":
                grid = tt_data.get("teacher_timetables", {}).get(name, {})
            elif cat == "class":
                grid = tt_data.get("student_timetables", {}).get(name, {})
            else:
                grid = tt_data.get("room_timetables", {}).get(name, {})

            break_slots = APP_STATE.institution.break_time_slots
            wed_settings = {
                "pe_label": APP_STATE.institution.pe_wednesday_label,
                "general_label": APP_STATE.institution.general_wednesday_label,
                "enabled": APP_STATE.institution.enable_wednesday_labels
            }
            sheet_html = EXPORT_ENGINE.render_single_sheet_html(
                cat, name, grid, APP_STATE.get_full_state()["institution"], break_slots,
                tt_data.get("generated_at", time.strftime("%Y-%m-%d")), wed_settings
            )
            full_page = f"""<!DOCTYPE html>
<html lang="ar" dir="rtl">
<head>
    <meta charset="UTF-8">
    <title>جدول حصص {name}</title>
    <style>{EXPORT_ENGINE.render_single_sheet_html.__globals__['get_base_css']()}</style>
</head>
<body>
    <div class="print-btn-bar">
        <button onclick="window.print()" class="btn-action">🖨️ طباعة الجدول (A4)</button>
        <button onclick="window.close()" class="btn-action" style="background:#475569;">✕ إغلاق</button>
    </div>
    {sheet_html}
</body>
</html>"""
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.end_headers()
            self.wfile.write(full_page.encode("utf-8"))
        else:
            super().do_GET()

    def do_POST(self):
        parsed = urllib.parse.urlparse(self.path)
        content_len = int(self.headers.get("Content-Length", 0))
        body = self.rfile.read(content_len)
        data = json.loads(body.decode("utf-8")) if body else {}

        if parsed.path == "/api/auth/activate_machine":
            email = data.get("email", "").strip()
            name = data.get("name", "").strip()
            res = APP_STATE.auth_engine.activate_machine(email, name)
            self.send_json_data(res)

        elif parsed.path == "/api/auth/login_google":
            # Direct machine activation or fallback
            res = {"success": False, "error": "يرجى استخدام نموذج تفعيل الترخيص الداخلي للتطبيق."}
            self.send_json_data(res)

        elif parsed.path == "/api/auth/dev_login":
            email = data.get("email", "").strip()
            name = data.get("name", "").strip()
            res = APP_STATE.auth_engine.activate_machine(email, name)
            self.send_json_data(res)

        elif parsed.path == "/api/auth/logout":
            APP_STATE.auth_engine.clear_session()
            self.send_json_data({"success": True})

        elif parsed.path == "/api/exams/generate_matrix":
            exam_config = data.get("config", {})
            total_candidates = data.get("total_candidates")
            available_rooms = data.get("available_rooms")
            capacity_per_room = data.get("capacity_per_room")
            if total_candidates is not None:
                try:
                    total_candidates = int(total_candidates)
                except Exception:
                    total_candidates = None
            if available_rooms is not None:
                try:
                    exam_config["available_rooms_count"] = int(available_rooms)
                except Exception:
                    pass
            if capacity_per_room is not None:
                try:
                    exam_config["capacity_per_room"] = int(capacity_per_room)
                except Exception:
                    pass
            assignments = APP_STATE.curriculum_engine.calculate_assignments()
            teachers_list = [{
                "name": a.teacher_name,
                "subject": a.subject
            } for a in assignments if a.teacher_name]
            engine = ExamManagementEngine(APP_STATE.institution, APP_STATE.structure)
            res = engine.build_invigilation_matrix(exam_config, teachers_list, total_candidates)
            self.send_json_data(res)

        elif parsed.path == "/api/surveillance/generate_report":
            day_base = data.get("day_base", "الاثنين")
            date_str = data.get("date_str", "")
            absent_teachers = data.get("absent_teachers", [])
            student_stats = data.get("student_stats", {})
            incidents = data.get("incidents", "")
            supervisor_name = data.get("supervisor_name", "")
            notes = data.get("notes", "")

            tt_data = APP_STATE.timetable_engine.get_cached_timetable() or {}
            engine = SurveillanceDailyReportEngine(tt_data, APP_STATE.institution)
            res = engine.generate_full_report(
                day_base=day_base,
                date_str=date_str,
                absent_teachers=absent_teachers,
                student_stats=student_stats,
                incidents=incidents,
                supervisor_name=supervisor_name,
                notes=notes
            )
            self.send_json_data(res)

        elif parsed.path == "/api/mobility/load_sample":
            try:
                res = mobility_engine.load_demo_sample_data(structure=APP_STATE.structure, institution=APP_STATE.institution)
            except Exception as e:
                print(f"[Mobility Demo Error]: {e}")
                res = mobility_engine.load_demo_sample_data()
            self.send_json_data(res)

        elif parsed.path == "/api/mobility/set_status":
            code = data.get("massar_code")
            status = data.get("status")
            ok = mobility_engine.set_student_override(code, status)
            self.send_json_data({"success": ok, "summary": mobility_engine.compute_mobility_analysis()})

        elif parsed.path == "/api/mobility/import_file":
            is_initial = data.get("is_initial", True)
            file_base64 = data.get("file_base64", "")
            if file_base64:
                import base64
                try:
                    file_bytes = base64.b64decode(file_base64)
                    res = mobility_engine.parse_massar_excel(file_bytes, is_initial=is_initial)
                    self.send_json_data({"success": True, "res": res, "summary": mobility_engine.compute_mobility_analysis()})
                except Exception as e:
                    self.send_json_data({"success": False, "error": f"خطأ أثناء معالجة الملف: {str(e)}"})
            else:
                self.send_json_data({"success": False, "error": "لم يتم إرسال بيانات ملف صالحة."})

        elif parsed.path == "/api/mobility/update_institution":
            if "institution" in data:
                mobility_engine.institution_info.update(data["institution"])
                mobility_engine.save_state()
            self.send_json_data({"success": True, "institution": mobility_engine.institution_info})

        elif parsed.path == "/api/open_in_browser":
            global LATEST_PRINT_DOC
            doc_html = data.get("html", "")
            doc_title = data.get("title", "وثيقة رسمية")
            LATEST_PRINT_DOC = {"title": doc_title, "html": doc_html}
            port = self.server.server_address[1]
            preview_url = f"http://127.0.0.1:{port}/print_preview?t={int(time.time())}"
            threading.Thread(target=lambda: webbrowser.open(preview_url), daemon=True).start()
            self.send_json_data({"success": True, "url": preview_url})

        elif parsed.path == "/api/admin/update_user":
            auth_info = APP_STATE.auth_engine.get_auth_status()
            if not auth_info.get("is_admin"):
                self.send_json_data({"success": False, "error": "غير مصرح لك بإجراء هذه العملية"})
                return
            
            target_email = data.get("email")
            status = data.get("status")
            extend_years = data.get("extend_years")
            new_expires_at = data.get("new_expires_at")
            sub_type = data.get("sub_type")
            res = APP_STATE.auth_engine.update_user_subscription(target_email, status, extend_years, new_expires_at, sub_type)
            self.send_json_data(res)

        elif parsed.path == "/api/admin/delete_user":
            auth_info = APP_STATE.auth_engine.get_auth_status()
            if not auth_info.get("is_admin"):
                self.send_json_data({"success": False, "error": "غير مصرح لك بإجراء هذه العملية"})
                return
            
            target_email = data.get("email")
            res = APP_STATE.auth_engine.delete_user(target_email)
            self.send_json_data(res)

        elif parsed.path == "/api/admin/publish_update":
            auth_info = APP_STATE.auth_engine.get_auth_status()
            if not auth_info.get("is_admin"):
                self.send_json_data({"success": False, "error": "غير مصرح لك بإجراء هذه العملية"})
                return
            new_version = data.get("new_version", "1.0.1")
            download_url = data.get("download_url", "")
            release_notes = data.get("release_notes", "")
            force_update = data.get("force_update", False)
            res = APP_STATE.auth_engine.publish_update(new_version, download_url, release_notes, force_update)
            self.send_json_data(res)

        elif parsed.path == "/api/admin/publish_live_patch":
            auth_info = APP_STATE.auth_engine.get_auth_status()
            if not (auth_info.get("is_admin") or auth_info.get("machine_id") == "RZZAK-F4A5-5734-7455"):
                self.send_json_data({"success": False, "error": "غير مصرح لك بنشر التحديثات البرمجية."})
                return

            new_version = str(data.get("new_version", "")).strip()
            release_notes = str(data.get("release_notes", "")).strip() or "تحسينات شاملة في الواجهة وتوليد الجداول."
            patch_type = data.get("patch_type", "live_patch")

            if not new_version:
                self.send_json_data({"success": False, "error": "يرجى تحديد رقم الإصدار الجديد."})
                return

            try:
                import subprocess
                # 1. Update config/auth_config.json version locally
                cfg_path = os.path.join(BASE_DIR, "config", "auth_config.json")
                if os.path.exists(cfg_path):
                    with open(cfg_path, "r", encoding="utf-8") as f:
                        cfg_json = json.load(f)
                    cfg_json.setdefault("app_info", {})["version"] = new_version
                    with open(cfg_path, "w", encoding="utf-8") as f:
                        json.dump(cfg_json, f, ensure_ascii=False, indent=2)

                # 2. Push git commit and push to origin main
                subprocess.run(["git", "add", "ui/", "src/", "config/", ".gitignore", "README.md"], capture_output=True, text=True, cwd=BASE_DIR)
                subprocess.run(["git", "commit", "-m", f"feat(ota): v{new_version} - {release_notes}"], capture_output=True, text=True, cwd=BASE_DIR)
                subprocess.run(["git", "push", "origin", "main"], capture_output=True, text=True, cwd=BASE_DIR)

                commit_sha = "main"
                try:
                    p_sha = subprocess.run(["git", "rev-parse", "HEAD"], capture_output=True, text=True, cwd=BASE_DIR)
                    if p_sha.returncode == 0 and p_sha.stdout.strip():
                        commit_sha = p_sha.stdout.strip()
                except Exception:
                    pass

                patch_url = f"https://raw.githubusercontent.com/abdellahrzzak/RzzakFet/{commit_sha}/ui/index.html"

                # 3. Publish to Firestore cloud metadata
                res = APP_STATE.auth_engine.publish_live_patch(
                    new_version=new_version,
                    release_notes=release_notes,
                    patch_type=patch_type,
                    patch_url=patch_url
                )

                # 4. Save local version to DATA_DIR so developer machine is registered on new version
                iv_file = os.path.join(DATA_DIR, "installed_version.json")
                with open(iv_file, "w", encoding="utf-8") as f:
                    json.dump({"version": new_version, "updated_at": time.strftime("%Y-%m-%dT%H:%M:%S")}, f, indent=2)

                self.send_json_data({
                    "success": True,
                    "version": new_version,
                    "data": res.get("data", {})
                })
            except Exception as e:
                self.send_json_data({"success": False, "error": str(e)})

        elif parsed.path == "/api/app/apply_live_patch":
            patch_url = data.get("patch_url") or "https://raw.githubusercontent.com/abdellahrzzak/RzzakFet/main/ui/index.html"
            new_version = str(data.get("new_version", "")).strip()

            try:
                fetch_url = patch_url
                if "?" in fetch_url:
                    fetch_url += f"&_t={int(time.time())}"
                else:
                    fetch_url += f"?_t={int(time.time())}"

                req = urllib.request.Request(fetch_url, headers={
                    "User-Agent": "RzzakFet-Desktop/2.0",
                    "Cache-Control": "no-cache"
                })
                with urllib.request.urlopen(req, timeout=12) as resp:
                    new_html = resp.read()

                if len(new_html) < 10000 or b"<!DOCTYPE html" not in new_html:
                    self.send_json_data({"success": False, "error": "ملف التحديث غير صالح أو لم يكتمل تحميله."})
                    return

                # Write to DATA_DIR/ui/index.html
                patch_dir = os.path.join(DATA_DIR, "ui")
                os.makedirs(patch_dir, exist_ok=True)
                patch_file = os.path.join(patch_dir, "index.html")
                with open(patch_file, "wb") as f:
                    f.write(new_html)

                # Attempt updating UI_DIR/index.html
                try:
                    ui_file = os.path.join(UI_DIR, "index.html")
                    with open(ui_file, "wb") as f:
                        f.write(new_html)
                except Exception:
                    pass

                # Attempt updating LOCALAPPDATA installed program
                try:
                    local_app_ui = os.path.join(os.environ.get("LOCALAPPDATA", ""), "Programs", "RzzakFet", "ui", "index.html")
                    if os.path.exists(os.path.dirname(local_app_ui)):
                        with open(local_app_ui, "wb") as f:
                            f.write(new_html)
                    internal_ui = os.path.join(os.environ.get("LOCALAPPDATA", ""), "Programs", "RzzakFet", "_internal", "ui", "index.html")
                    if os.path.exists(os.path.dirname(internal_ui)):
                        with open(internal_ui, "wb") as f:
                            f.write(new_html)
                except Exception:
                    pass

                # Record installed version in DATA_DIR
                if new_version:
                    iv_file = os.path.join(DATA_DIR, "installed_version.json")
                    with open(iv_file, "w", encoding="utf-8") as f:
                        json.dump({"version": new_version, "updated_at": time.strftime("%Y-%m-%dT%H:%M:%S")}, f, indent=2)

                self.send_json_data({"success": True, "version": new_version})
            except Exception as e:
                self.send_json_data({"success": False, "error": f"فشل تثبيت التحديث: {str(e)}"})

        elif parsed.path == "/api/export_to_desktop":
            tt_data = APP_STATE.timetable_engine.get_cached_timetable()
            if not tt_data:
                assignments = APP_STATE.curriculum_engine.calculate_assignments()
                rooms = APP_STATE.room_engine.allocate_rooms(assignments)
                xml_data = APP_STATE.fet_generator.generate_xml(assignments, rooms)
                tt_data = APP_STATE.timetable_engine.generate_timetable(xml_data, APP_STATE.institution, rooms)
            
            res = EXPORT_ENGINE.export_all(tt_data, APP_STATE.get_full_state())
            # Auto-open the folder
            EXPORT_ENGINE.open_desktop_folder()

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            self.wfile.write(json.dumps(res, ensure_ascii=False).encode("utf-8"))

        elif parsed.path == "/api/start_progressive_generation":
            res = PROGRESSIVE_MANAGER.start_generation()
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            self.wfile.write(json.dumps(res, ensure_ascii=False).encode("utf-8"))

        elif parsed.path == "/api/stop_progressive_generation":
            res = PROGRESSIVE_MANAGER.stop_generation()
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            self.wfile.write(json.dumps(res, ensure_ascii=False).encode("utf-8"))

        elif parsed.path == "/api/launch_fet_gui":
            assignments = APP_STATE.curriculum_engine.calculate_assignments()
            rooms = APP_STATE.room_engine.allocate_rooms(assignments)
            xml_data = APP_STATE.fet_generator.generate_xml(assignments, rooms)
            res = APP_STATE.native_bridge.launch_fet_gui(xml_data, APP_STATE.institution.institution_name)

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            self.wfile.write(json.dumps(res, ensure_ascii=False).encode("utf-8"))

        elif parsed.path == "/api/sync_fet_results":
            latest_xml = APP_STATE.native_bridge.find_latest_generated_xml()
            if latest_xml and os.path.exists(latest_xml):
                assignments = APP_STATE.curriculum_engine.calculate_assignments()
                rooms = APP_STATE.room_engine.allocate_rooms(assignments)
                try:
                    with open(latest_xml, "r", encoding="utf-8") as f:
                        xml_content = f.read()
                    res = APP_STATE.timetable_engine.generate_timetable(xml_content, APP_STATE.institution, rooms)
                    res["synced_from"] = latest_xml
                except Exception as e:
                    res = {"success": False, "error": f"خطأ أثناء قراءة ملف المخرجات: {str(e)}"}
            else:
                res = {"success": False, "error": "لم يتم العثور على أي ملف مخرجات حديث من FET في مجلد timetables."}

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            self.wfile.write(json.dumps(res, ensure_ascii=False).encode("utf-8"))

        elif parsed.path == "/api/generate_timetable":
            assignments = APP_STATE.curriculum_engine.calculate_assignments()
            rooms = APP_STATE.room_engine.allocate_rooms(assignments)
            xml_data = APP_STATE.fet_generator.generate_xml(assignments, rooms)
            res = APP_STATE.timetable_engine.generate_timetable(xml_data, APP_STATE.institution, rooms)

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            self.wfile.write(json.dumps(res, ensure_ascii=False).encode("utf-8"))

        elif parsed.path == "/api/stop_generation":
            APP_STATE.timetable_engine.stop_generation()
            PROGRESSIVE_MANAGER.stop_generation()
            res = {"success": True, "message": "تم إيقاف عملية الإنتاج بنجاح."}

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            self.wfile.write(json.dumps(res, ensure_ascii=False).encode("utf-8"))

        elif parsed.path == "/api/export_fet_file":
            try:
                assignments = APP_STATE.curriculum_engine.calculate_assignments()
                rooms = APP_STATE.room_engine.allocate_rooms(assignments)
                xml_data = APP_STATE.fet_generator.generate_xml(assignments, rooms)
                
                gresa = APP_STATE.institution.gresa_code or "institution"
                filename = f"RzzakFet_{gresa}.fet"
                
                desktop_rzzak_dir = os.path.join(os.path.expanduser("~"), "Desktop", "RzzakFet")
                os.makedirs(desktop_rzzak_dir, exist_ok=True)
                desktop_file_path = os.path.join(desktop_rzzak_dir, filename)
                
                with open(desktop_file_path, "w", encoding="utf-8") as f:
                    f.write(xml_data)
                    
                data_fet_path = os.path.join(DATA_DIR, filename)
                with open(data_fet_path, "w", encoding="utf-8") as f:
                    f.write(xml_data)
                
                res = {
                    "success": True,
                    "filename": filename,
                    "xml_content": xml_data,
                    "desktop_path": desktop_file_path,
                    "message": "تم تصدير وحفظ ملف FET بنجاح."
                }
            except Exception as e:
                res = {"success": False, "error": str(e)}

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            self.wfile.write(json.dumps(res, ensure_ascii=False).encode("utf-8"))

        elif parsed.path == "/api/update_institution":
            inst = APP_STATE.institution
            inst.institution_name = data.get("name", inst.institution_name)
            inst.province = data.get("province", inst.province)
            inst.academy = data.get("academy", inst.academy)
            inst.academic_year = data.get("academic_year", inst.academic_year)
            inst.gresa_code = data.get("gresa", inst.gresa_code)
            inst.principal_name = data.get("principal", inst.principal_name)
            inst.general_rooms_count = int(data.get("general_rooms", inst.general_rooms_count))
            inst.svt_labs_count = int(data.get("svt_labs", inst.svt_labs_count))
            inst.pc_labs_count = int(data.get("pc_labs", inst.pc_labs_count))
            inst.multimedia_rooms_count = int(data.get("multimedia_rooms", inst.multimedia_rooms_count))
            inst.sports_fields_count = int(data.get("sports_fields", inst.sports_fields_count))
            if "morning_start_hour" in data:
                inst.morning_start_hour = str(data["morning_start_hour"]).strip()
            if "morning_period_times" in data and isinstance(data["morning_period_times"], list):
                inst.morning_period_times = [str(x).strip() for x in data["morning_period_times"] if str(x).strip()]
            if "afternoon_start_hour" in data:
                inst.afternoon_start_hour = str(data["afternoon_start_hour"]).strip()
            if "afternoon_period_times" in data and isinstance(data["afternoon_period_times"], list):
                inst.afternoon_period_times = [str(x).strip() for x in data["afternoon_period_times"] if str(x).strip()]
            APP_STATE.save_to_disk()
            
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = json.dumps(APP_STATE.get_full_state(), ensure_ascii=False)
            self.wfile.write(res.encode("utf-8"))

        elif parsed.path == "/api/update_room_names":
            names_map = data.get("names", {})
            APP_STATE.institution.room_custom_names = names_map
            APP_STATE.save_to_disk()

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = json.dumps(APP_STATE.get_full_state(), ensure_ascii=False)
            self.wfile.write(res.encode("utf-8"))

        elif parsed.path in ["/api/save_teacher_names", "/api/update_teacher_names"]:
            names_map = data.get("names", {})
            clean_map = {str(k).strip(): str(v).strip() for k, v in names_map.items() if str(v).strip()}
            APP_STATE.institution.teacher_custom_names = clean_map
            APP_STATE.save_to_disk()
            APP_STATE.curriculum_engine.calculate_assignments()
            APP_STATE.remap_timetable_teachers()

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = json.dumps(APP_STATE.get_full_state(), ensure_ascii=False)
            self.wfile.write(res.encode("utf-8"))

        elif parsed.path == "/api/export_teachers_excel":
            file_bytes, dest_path = export_teachers_to_excel(APP_STATE)
            file_b64 = base64.b64encode(file_bytes).decode("ascii")
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = json.dumps({
                "success": True,
                "filename": "لائحة_أسماء_الأساتذة_RzzakFet.xlsx",
                "desktop_path": dest_path,
                "file_base64": file_b64
            }, ensure_ascii=False)
            self.wfile.write(res.encode("utf-8"))

        elif parsed.path == "/api/import_teachers_excel":
            file_b64 = data.get("content", "")
            filename = data.get("filename", "teachers.xlsx")
            file_bytes = base64.b64decode(file_b64) if file_b64 else b""
            result = import_teachers_from_file(APP_STATE, file_bytes, filename)
            result["full_state"] = APP_STATE.get_full_state()
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = json.dumps(result, ensure_ascii=False)
            self.wfile.write(res.encode("utf-8"))

        elif parsed.path == "/api/toggle_float_all":
            current = APP_STATE.institution.float_all_general_teachers
            APP_STATE.institution.float_all_general_teachers = not current
            APP_STATE.save_to_disk()

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = json.dumps(APP_STATE.get_full_state(), ensure_ascii=False)
            self.wfile.write(res.encode("utf-8"))

        elif parsed.path == "/api/add_custom_room_type":
            name = data.get("name", "قاعة مخصصة")
            subj = data.get("subject", "عامة")
            r_type = data.get("room_type", "قاعة مخصصة")
            count = int(data.get("count", 1))
            crt_id = str(uuid.uuid4())[:8]

            APP_STATE.institution.custom_room_types.append({
                "id": crt_id, "name": name, "subject": subj,
                "room_type": r_type, "count": max(1, count)
            })
            APP_STATE.save_to_disk()

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = json.dumps(APP_STATE.get_full_state(), ensure_ascii=False)
            self.wfile.write(res.encode("utf-8"))

        elif parsed.path == "/api/update_custom_room_type_count":
            crt_id = data.get("id")
            new_count = int(data.get("count", 1))
            for crt in APP_STATE.institution.custom_room_types:
                if crt.get("id") == crt_id:
                    crt["count"] = max(1, new_count)
                    break
            APP_STATE.save_to_disk()

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = json.dumps(APP_STATE.get_full_state(), ensure_ascii=False)
            self.wfile.write(res.encode("utf-8"))

        elif parsed.path == "/api/delete_custom_room_type":
            crt_id = data.get("id")
            APP_STATE.institution.custom_room_types = [
                crt for crt in APP_STATE.institution.custom_room_types if crt.get("id") != crt_id
            ]
            APP_STATE.save_to_disk()

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = json.dumps(APP_STATE.get_full_state(), ensure_ascii=False)
            self.wfile.write(res.encode("utf-8"))

        elif parsed.path == "/api/update_room_teacher":
            room_name = data.get("room_name")
            session = data.get("session")
            teacher_name = data.get("teacher_name")
            if room_name and session:
                APP_STATE.room_engine.set_teacher(room_name, session, teacher_name)
                APP_STATE.save_to_disk()

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = json.dumps(APP_STATE.get_full_state(), ensure_ascii=False)
            self.wfile.write(res.encode("utf-8"))

        elif parsed.path == "/api/reallocate_rooms":
            clear_overrides = data.get("clear_overrides", False)
            if clear_overrides:
                APP_STATE.room_engine.clear_all_overrides()
            APP_STATE.save_to_disk()

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = json.dumps(APP_STATE.get_full_state(), ensure_ascii=False)
            self.wfile.write(res.encode("utf-8"))

        elif parsed.path == "/api/update_structure":
            APP_STATE.structure.classes_1_apic = int(data.get("c1", APP_STATE.structure.classes_1_apic))
            APP_STATE.structure.classes_2_apic = int(data.get("c2", APP_STATE.structure.classes_2_apic))
            APP_STATE.structure.classes_3_apic = int(data.get("c3", APP_STATE.structure.classes_3_apic))
            if "s1" in data:
                APP_STATE.structure.students_1_apic = int(data.get("s1", APP_STATE.structure.students_1_apic))
            if "s2" in data:
                APP_STATE.structure.students_2_apic = int(data.get("s2", APP_STATE.structure.students_2_apic))
            if "s3" in data:
                APP_STATE.structure.students_3_apic = int(data.get("s3", APP_STATE.structure.students_3_apic))
            APP_STATE.save_to_disk()
            
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = json.dumps(APP_STATE.get_full_state(), ensure_ascii=False)
            self.wfile.write(res.encode("utf-8"))

        elif parsed.path == "/api/toggle_subject":
            subj_name = data.get("name")
            for cfg in APP_STATE.curriculum_engine.non_generalized_configs:
                if cfg.name == subj_name:
                    cfg.is_active = not cfg.is_active
                    break
            APP_STATE.save_to_disk()

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = json.dumps(APP_STATE.get_full_state(), ensure_ascii=False)
            self.wfile.write(res.encode("utf-8"))

        elif parsed.path == "/api/update_break_times":
            slots = data.get("slots", [])
            APP_STATE.institution.break_time_slots = slots
            
            for c in APP_STATE.fet_generator.all_available_constraints:
                if c["id"] in ("t_all_1", "tc_base_2"):
                    c["param_val"] = len(slots)
                    break
            APP_STATE.save_to_disk()

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = json.dumps(APP_STATE.get_full_state(), ensure_ascii=False)
            self.wfile.write(res.encode("utf-8"))

        elif parsed.path == "/api/update_constraint":
            c_id = data.get("id")
            is_active = data.get("is_active")
            is_primary = data.get("is_primary")
            weight = data.get("weight")
            param_val = data.get("param_val")
            applies_to = data.get("applies_to")
            selected_targets = data.get("selected_targets")

            for c in APP_STATE.fet_generator.all_available_constraints:
                if c["id"] == c_id:
                    if is_active is not None:
                        c["is_active"] = bool(is_active)
                    if is_primary is not None:
                        c["is_primary"] = bool(is_primary)
                    if weight is not None:
                        c["def_weight"] = float(weight)
                    if param_val is not None and c.get("param_name"):
                        try:
                            c["param_val"] = int(param_val)
                        except (ValueError, TypeError):
                            c["param_val"] = param_val
                    if applies_to is not None:
                        c["applies_to"] = str(applies_to)
                    if selected_targets is not None and isinstance(selected_targets, list):
                        c["selected_targets"] = [str(x) for x in selected_targets]
                    break
            APP_STATE.save_to_disk()

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = json.dumps(APP_STATE.get_full_state(), ensure_ascii=False)
            self.wfile.write(res.encode("utf-8"))

        elif parsed.path == "/api/update_constraint_target":
            c_id = data.get("id")
            applies_to = data.get("applies_to", "all")
            selected_targets = data.get("selected_targets", [])
            for c in APP_STATE.fet_generator.all_available_constraints:
                if c["id"] == c_id:
                    c["applies_to"] = str(applies_to)
                    if isinstance(selected_targets, list):
                        c["selected_targets"] = [str(x) for x in selected_targets]
                    break
            APP_STATE.save_to_disk()

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = json.dumps(APP_STATE.get_full_state(), ensure_ascii=False)
            self.wfile.write(res.encode("utf-8"))

        elif parsed.path == "/api/reorder_constraints":
            new_order = data.get("order", [])
            if new_order and isinstance(new_order, list):
                c_map = {c["id"]: c for c in APP_STATE.fet_generator.all_available_constraints}
                ordered_primary = []
                for cid in new_order:
                    if cid in c_map:
                        c_map[cid]["is_primary"] = True
                        ordered_primary.append(c_map[cid])
                # Remaining constraints keep their relative positions
                remaining = [c for c in APP_STATE.fet_generator.all_available_constraints if c["id"] not in set(new_order)]
                APP_STATE.fet_generator.all_available_constraints = ordered_primary + remaining
                APP_STATE.save_to_disk()

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = json.dumps(APP_STATE.get_full_state(), ensure_ascii=False)
            self.wfile.write(res.encode("utf-8"))

        elif parsed.path == "/api/delete_constraint":
            c_id = data.get("id")
            for c in APP_STATE.fet_generator.all_available_constraints:
                if c["id"] == c_id:
                    c["is_primary"] = False
                    c["is_active"] = False
                    break
            APP_STATE.save_to_disk()

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = json.dumps(APP_STATE.get_full_state(), ensure_ascii=False)
            self.wfile.write(res.encode("utf-8"))

        elif parsed.path == "/api/add_constraint_from_library":
            c_id = data.get("id")
            target = None
            idx = -1
            all_c = APP_STATE.fet_generator.all_available_constraints
            for i, c in enumerate(all_c):
                if c["id"] == c_id:
                    target = c
                    idx = i
                    break
            if target:
                target["is_primary"] = True
                target["is_active"] = True
                all_c.pop(idx)
                # Insert at the end of the primary list
                last_primary_idx = 0
                for i, c in enumerate(all_c):
                    if c.get("is_primary", False):
                        last_primary_idx = i + 1
                all_c.insert(last_primary_idx, target)
                APP_STATE.save_to_disk()

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = json.dumps(APP_STATE.get_full_state(), ensure_ascii=False)
            self.wfile.write(res.encode("utf-8"))

        elif parsed.path == "/api/reset_default_constraints":
            APP_STATE.fet_generator.reset_to_default_constraints()
            APP_STATE.save_to_disk()

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = json.dumps(APP_STATE.get_full_state(), ensure_ascii=False)
            self.wfile.write(res.encode("utf-8"))

        elif parsed.path in ("/api/apply_core_9_constraints", "/api/apply_core_11_constraints"):
            toggle = data.get("toggle", False)
            core_codes = [
                "ConstraintTeachersMaxHoursDaily",
                "ConstraintTeachersMaxGapsPerDay",
                "ConstraintTeachersMaxContinuousHours",
                "ConstraintStudentsMaxGapsPerDay",
                "ConstraintStudentsMaxHoursDaily",
                "ConstraintMinDaysBetweenActivities",
                "ConstraintBreakTimes",
                "ConstraintTeacherHomeRoom",
                "ConstraintSubjectPreferredRooms"
            ]

            # Check if all core are currently active
            all_active = all(
                c.get("is_active", False) 
                for c in APP_STATE.fet_generator.all_available_constraints 
                if c.get("code") in core_codes and c.get("is_primary", False)
            )

            new_target_state = False if (toggle and all_active) else True

            for c in APP_STATE.fet_generator.all_available_constraints:
                code = c.get("code")
                if code in core_codes:
                    c["is_primary"] = True
                    c["is_active"] = new_target_state
                    c["def_weight"] = 100.0
                    if code in ["ConstraintTeachersMaxHoursDaily", "ConstraintStudentsMaxHoursDaily"]:
                        c["param_val"] = 6
                    elif code == "ConstraintTeachersMaxContinuousHours":
                        c["param_val"] = 4
                    elif code in ["ConstraintTeachersMaxGapsPerDay", "ConstraintStudentsMaxGapsPerDay"]:
                        c["param_val"] = 0
                    elif code == "ConstraintMinDaysBetweenActivities":
                        c["param_val"] = 1
                    elif code == "ConstraintBreakTimes":
                        c["param_val"] = len(APP_STATE.institution.break_time_slots)

            APP_STATE.save_to_disk()

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = json.dumps(APP_STATE.get_full_state(), ensure_ascii=False)
            self.wfile.write(res.encode("utf-8"))

        elif parsed.path == "/api/update_subject_activity_constraint":
            subject = data.get("subject")
            if subject:
                cur_map = APP_STATE.institution.subject_activity_constraints
                if subject not in cur_map:
                    cur_map[subject] = {}
                
                if "min_days" in data:
                    cur_map[subject]["min_days"] = int(data["min_days"])
                if "max_days" in data:
                    cur_map[subject]["max_days"] = int(data["max_days"]) if data["max_days"] is not None and str(data["max_days"]).strip() != "" else None
                if "weight" in data:
                    cur_map[subject]["weight"] = float(data["weight"])
                if "is_active" in data:
                    cur_map[subject]["is_active"] = bool(data["is_active"])
                if "consecutive_if_same_day" in data:
                    cur_map[subject]["consecutive_if_same_day"] = bool(data["consecutive_if_same_day"])
                if "forbid_first_pm_slot" in data:
                    cur_map[subject]["forbid_first_pm_slot"] = bool(data["forbid_first_pm_slot"])

                APP_STATE.save_to_disk()

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = json.dumps(APP_STATE.get_full_state(), ensure_ascii=False)
            self.wfile.write(res.encode("utf-8"))

        elif parsed.path == "/api/update_subject_splits":
            subject = data.get("subject")
            mode = data.get("mode")
            mode_l1 = data.get("mode_l1")
            mode_l2 = data.get("mode_l2")
            mode_l3 = data.get("mode_l3")
            desc = data.get("desc")

            if subject and mode:
                splits = APP_STATE.institution.subject_split_modes
                if subject not in splits:
                    splits[subject] = {}
                splits[subject]["mode"] = str(mode)
                if mode_l1:
                    splits[subject]["mode_l1"] = str(mode_l1)
                if mode_l2:
                    splits[subject]["mode_l2"] = str(mode_l2)
                if mode_l3:
                    splits[subject]["mode_l3"] = str(mode_l3)
                if desc:
                    splits[subject]["desc"] = str(desc)
                APP_STATE.save_to_disk()

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = json.dumps(APP_STATE.get_full_state(), ensure_ascii=False)
            self.wfile.write(res.encode("utf-8"))

        elif parsed.path == "/api/reset_subject_splits":
            # Reset to defaults
            APP_STATE.institution.subject_split_modes = {
                "التربية البدنية": {"mode": "2", "desc": "حصة واحدة من ساعتين بالملاعب (النمط الرسمي)"},
                "اللغة العربية": {"mode": "1+1+1+1", "desc": "4 حصص فردية على 4 أيام (المذكرة 43)"},
                "اللغة الفرنسية": {"mode": "1+1+1+1", "desc": "4 حصص فردية على 4 أيام (المذكرة 43)"},
                "الرياضيات": {"mode": "1+1+1+1+1", "mode_l1": "1+1+1+1+1", "mode_l2": "1+1+1+1", "mode_l3": "1+1+1+1+1", "desc": "حصص فردية يومية (5 للأولى والثالثة، 4 للثانية)"},
                "الاجتماعيات": {"mode": "1+1+1", "desc": "3 حصص فردية على 3 أيام مختلفة"},
                "التربية الإسلامية": {"mode": "1+1", "desc": "حصتان فرديتان على يومين متباعدين"},
                "علوم الحياة و الأرض": {"mode": "2", "mode_l3": "2+1", "desc": "ساعتان متصلتان بالمختبر (2+1 للثالثة إعدادي)"},
                "الكيمياء و الفيزياء": {"mode": "2", "mode_l3": "2+1", "desc": "ساعتان متصلتان بالمختبر (2+1 للثالثة إعدادي)"}
            }
            APP_STATE.save_to_disk()

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = json.dumps(APP_STATE.get_full_state(), ensure_ascii=False)
            self.wfile.write(res.encode("utf-8"))

        elif parsed.path == "/api/update_wednesday_settings":
            if "pe_label" in data:
                APP_STATE.institution.pe_wednesday_label = str(data["pe_label"]).strip()
            if "general_label" in data:
                APP_STATE.institution.general_wednesday_label = str(data["general_label"]).strip()
            if "enabled" in data:
                APP_STATE.institution.enable_wednesday_labels = bool(data["enabled"])
            APP_STATE.save_to_disk()

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = json.dumps(APP_STATE.get_full_state(), ensure_ascii=False)
            self.wfile.write(res.encode("utf-8"))

        elif parsed.path == "/api/update_teacher_count":
            subj = data.get("subject")
            count = data.get("count")
            if subj and count is not None:
                try:
                    APP_STATE.institution.teacher_counts_by_subject[str(subj).strip()] = max(0, int(count))
                    APP_STATE.save_to_disk()
                except Exception as e:
                    print("Error updating teacher count:", e)

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = json.dumps(APP_STATE.get_full_state(), ensure_ascii=False)
            self.wfile.write(res.encode("utf-8"))

        elif parsed.path == "/api/reset_teacher_counts":
            quotas = APP_STATE.curriculum_engine.get_subject_quotas()
            APP_STATE.institution.teacher_counts_by_subject = {
                q.name: q.required_teachers for q in quotas if q.total_hours > 0
            }
            APP_STATE.save_to_disk()

            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            res = json.dumps(APP_STATE.get_full_state(), ensure_ascii=False)
            self.wfile.write(res.encode("utf-8"))

        else:
            self.send_error(404)

def run_server(port=8765):
    server_address = ('', port)
    httpd = http.server.HTTPServer(server_address, RzzakFetHandler)
    print(f"RzzakFet Server running on port {port}...")
    httpd.serve_forever()

if __name__ == "__main__":
    run_server()
