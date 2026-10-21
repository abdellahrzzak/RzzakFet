# -*- coding: utf-8 -*-
import os
import subprocess
import json
import time
from typing import Dict, List, Any, Optional
from concurrent.futures import ThreadPoolExecutor
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DAYS_ORDER = [
    {"name": "الإثنين", "m": "الاثنين ص", "e": "الاثنين م"},
    {"name": "الثلاثاء", "m": "الثلاثاء ص", "e": "الثلاثاء م"},
    {"name": "الأربعاء", "m": "الاربعاء ص", "e": "الاربعاء م"},
    {"name": "الخميس", "m": "الخميس ص", "e": "الخميس م"},
    {"name": "الجمعة", "m": "الجمعة ص", "e": "الجمعة م"},
    {"name": "السبت", "m": "السبت ص", "e": "السبت م"}
]

HOURS_4 = ["ح 1", "ح 2", "ح 3", "ح 4"]

SUBJECT_COLORS = {
    "عربي": {"bg": "#ffe4e6", "border": "#f43f5e", "text": "#9f1239", "print_bg": "#f43f5e", "print_text": "#ffffff"},
    "رياض": {"bg": "#ecfccb", "border": "#84cc16", "text": "#3f6212", "print_bg": "#65a30d", "print_text": "#ffffff"},
    "اجتماع": {"bg": "#ede9fe", "border": "#8b5cf6", "text": "#5b21b6", "print_bg": "#7c3aed", "print_text": "#ffffff"},
    "فرنس": {"bg": "#fef3c7", "border": "#f59e0b", "text": "#92400e", "print_bg": "#d97706", "print_text": "#ffffff"},
    "إسلام": {"bg": "#e0f2fe", "border": "#0284c7", "text": "#075985", "print_bg": "#0284c7", "print_text": "#ffffff"},
    "حياة": {"bg": "#fef9c3", "border": "#ca8a04", "text": "#854d0e", "print_bg": "#b45309", "print_text": "#ffffff"},
    "SVT": {"bg": "#fef9c3", "border": "#ca8a04", "text": "#854d0e", "print_bg": "#b45309", "print_text": "#ffffff"},
    "فيز": {"bg": "#dcfce7", "border": "#22c55e", "text": "#166534", "print_bg": "#16a34a", "print_text": "#ffffff"},
    "PC": {"bg": "#dcfce7", "border": "#22c55e", "text": "#166534", "print_bg": "#16a34a", "print_text": "#ffffff"},
    "بدن": {"bg": "#ffe4e6", "border": "#fb7185", "text": "#9f1239", "print_bg": "#e11d48", "print_text": "#ffffff"},
    "EPS": {"bg": "#ffe4e6", "border": "#fb7185", "text": "#9f1239", "print_bg": "#e11d48", "print_text": "#ffffff"},
    "معلوم": {"bg": "#cffafe", "border": "#06b6d4", "text": "#155e75", "print_bg": "#0891b2", "print_text": "#ffffff"},
    "إنجل": {"bg": "#ffedd5", "border": "#f97316", "text": "#9a3412", "print_bg": "#ea580c", "print_text": "#ffffff"}
}

def get_subj_style(subject: str) -> Dict[str, str]:
    if not subject:
        return {"bg": "#f8fafc", "border": "#cbd5e1", "text": "#0f172a", "print_bg": "#475569", "print_text": "#ffffff"}
    for k, v in SUBJECT_COLORS.items():
        if k in subject:
            return v
    return {"bg": "#f1f5f9", "border": "#94a3b8", "text": "#1e293b", "print_bg": "#475569", "print_text": "#ffffff"}

import re

LEVEL_1_PALETTES = ["#2563eb", "#0284c7", "#0891b2", "#0d9488", "#3b82f6", "#06b6d4", "#14b8a6", "#0369a1", "#0e7490", "#0f766e", "#1d4ed8", "#075985"]
LEVEL_2_PALETTES = ["#059669", "#16a34a", "#65a30d", "#ca8a04", "#10b981", "#22c55e", "#84cc16", "#eab308", "#047857", "#15803d", "#4d7c0f", "#a16207"]
LEVEL_3_PALETTES = ["#ea580c", "#e11d48", "#9333ea", "#7c3aed", "#c026d3", "#db2777", "#f97316", "#f43f5e", "#a855f7", "#8b5cf6", "#d946ef", "#ec4899"]
FALLBACK_PALETTES = ["#4f46e5", "#d97706", "#dc2626", "#475569", "#6366f1", "#b45309"]

def get_class_style(class_name: str) -> Dict[str, str]:
    if not class_name:
        return {"print_bg": "#475569", "print_text": "#ffffff"}
    s = str(class_name).strip()
    m_level = re.search(r'([1-3])', s)
    m_sec = re.search(r'(\d+)$', s)
    sec_num = int(m_sec.group(1)) if m_sec else 1
    lvl_num = int(m_level.group(1)) if m_level else 0

    if lvl_num == 1:
        bg = LEVEL_1_PALETTES[(sec_num - 1) % len(LEVEL_1_PALETTES)]
    elif lvl_num == 2:
        bg = LEVEL_2_PALETTES[(sec_num - 1) % len(LEVEL_2_PALETTES)]
    elif lvl_num == 3:
        bg = LEVEL_3_PALETTES[(sec_num - 1) % len(LEVEL_3_PALETTES)]
    else:
        bg = FALLBACK_PALETTES[sec_num % len(FALLBACK_PALETTES)]
    return {"print_bg": bg, "print_text": "#ffffff"}

def abbreviate_class_name(s: str) -> str:
    if not s:
        return ""
    s = str(s).strip()
    m = re.match(r'^([1-3])\s*[-_]?(?:APIC|ASC|AC|apic|asc|ac)?[-\s_]*(\d+)$', s, re.IGNORECASE)
    if m:
        return f"{m.group(1)}/{m.group(2)}"
    m2 = re.match(r'^([1-3])[A-Za-z]+(\d+)$', s)
    if m2:
        return f"{m2.group(1)}/{m2.group(2)}"
    return s

def abbreviate_room_name(s: str) -> str:
    if not s:
        return ""
    s = str(s).strip()
    if len(s) <= 4:
        return s
    m_num = re.search(r'(\d+)', s)
    num = m_num.group(1) if m_num else ''
    
    # Physics / Chemistry (العلوم الفيزيائية)
    if 'فيزيائ' in s or 'فيزياء' in s or 'PC' in s.upper():
        return f'PC{num}' if num else 'PC'
    # SVT (علوم الحياة والأرض)
    if 'حياة' in s or 'أرض' in s or 'SVT' in s.upper():
        return f'SVT{num}' if num else 'SVT'
    # Informatics
    if 'معلوم' in s or 'إعلام' in s or 'INFO' in s.upper():
        return f'INF{num}' if num else 'INFO'
    # Sport / EPS
    if 'بدن' in s or 'رياض' in s or 'EPS' in s.upper() or 'TERRAIN' in s.upper():
        return f'EPS{num}' if num else 'EPS'
    # General rooms
    if 'عامة' in s or 'GH' in s.upper() or 'S-' in s.upper() or 'قاعة' in s or 'S ' in s:
        return f'ق{num}' if num else 'ق'
    if num:
        return f'ق{num}'
    return s[:4]

def split_class_name_for_single_cell(s: str):
    if not s:
        return "", ""
    s = str(s).strip()
    parts = re.split(r'[-_\s]+', s)
    if len(parts) >= 3:
        if len(parts) >= 4:
            p1 = "-".join(parts[:-2])
            p2 = "-".join(parts[-2:])
        else:
            p1 = "-".join(parts[:-1])
            p2 = parts[-1]
        return p1, p2
    elif len(parts) == 2:
        return parts[0], parts[1]
    
    m = re.match(r'^([0-9]?[A-Za-z]+)[-_]?([0-9]+)$', s)
    if m:
        return m.group(1), m.group(2)
        
    return s, ""

def format_master_cell_class(stds: str, rm: str, span: int = 1) -> str:
    if not stds:
        return ""
    stds = str(stds).strip()
    short_rm = abbreviate_room_name(rm)
    rm_html = f'<div class="r-code">{short_rm}</div>' if short_rm else ""

    if span >= 2:
        return f'<div class="c-code full-span" title="{stds}">{stds}</div>{rm_html}'
    
    p1, p2 = split_class_name_for_single_cell(stds)
    if p2:
        return f'<div class="c-code single-span" title="{stds}"><span class="c-line1">{p1}</span><span class="c-line2">{p2}</span></div>{rm_html}'
    else:
        return f'<div class="c-code single-span" title="{stds}"><span class="c-line1">{p1}</span></div>{rm_html}'



def get_base_css() -> str:
    return """
    @import url('https://fonts.googleapis.com/css2?family=Cairo:wght@400;600;700;800;900&display=swap');
    
    * {
        box-sizing: border-box;
        font-family: 'Cairo', -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
        margin: 0;
        padding: 0;
    }
    body {
        background-color: #f8fafc;
        color: #0f172a;
        direction: rtl;
        padding: 15px;
    }
    .sheet {
        background: #ffffff;
        border: 2px solid #0f172a;
        border-radius: 8px;
        padding: 12px 16px;
        margin: 0 auto 20px auto;
        max-width: 1100px;
        box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.1);
        page-break-inside: avoid;
        page-break-after: always;
    }
    .sheet:last-child {
        page-break-after: auto;
    }
    .official-header {
        display: flex;
        justify-content: space-between;
        align-items: center;
        border-bottom: 2px solid #0f172a;
        padding-bottom: 6px;
        margin-bottom: 8px;
    }
    .header-box {
        font-size: 10.5pt;
        line-height: 1.35;
    }
    .header-center {
        text-align: center;
    }
    .main-title {
        border: 2px solid #0f172a;
        padding: 4px 16px;
        border-radius: 6px;
        font-size: 12pt;
        font-weight: 900;
        background: #f1f5f9;
        display: inline-block;
        margin-bottom: 3px;
    }
    .sub-hierarchy {
        display: flex;
        justify-content: space-between;
        background: #f1f5f9;
        border: 1px solid #cbd5e1;
        border-radius: 4px;
        padding: 4px 10px;
        font-size: 9.5pt;
        font-weight: 700;
        margin-bottom: 8px;
    }
    .grid-table {
        width: 100%;
        border-collapse: collapse;
        table-layout: fixed;
        border: 2px solid #0f172a;
        margin-bottom: 8px;
    }
    .grid-table th {
        background: #1e293b;
        color: #ffffff;
        font-size: 9.5pt;
        font-weight: 800;
        padding: 5px 2px;
        border: 1px solid #0f172a;
        text-align: center;
    }
    .grid-table td {
        border: 1.5px solid #334155;
        text-align: center;
        vertical-align: middle;
        padding: 4px 2px;
        height: 52px;
    }
    .day-cell {
        background: #e2e8f0;
        font-weight: 900;
        font-size: 11pt;
        color: #0f172a;
        width: 9%;
        border: 1.5px solid #0f172a !important;
    }
    .subj-cell {
        font-weight: 800;
        font-size: 10pt;
        line-height: 1.2;
    }
    .sub-info {
        font-size: 8.5pt;
        font-weight: 600;
        margin-top: 2px;
    }
    .cell-empty {
        background: repeating-linear-gradient(45deg, #cbd5e1 0, #cbd5e1 1.5px, #f8fafc 1.5px, #f8fafc 7px);
    }
    .cell-break {
        background: repeating-linear-gradient(-45deg, #94a3b8 0, #94a3b8 1.5px, #e2e8f0 1.5px, #e2e8f0 7px);
    }
    .cell-ass {
        background: #1e3a8a;
        color: #ffffff;
        font-weight: 800;
        font-size: 9pt;
        padding: 2px;
    }
    .cell-support {
        background: #064e3b;
        color: #ffffff;
        font-weight: 800;
        font-size: 9pt;
        padding: 2px;
    }
    .assigned-table {
        width: 100%;
        border-collapse: collapse;
        border: 1.5px solid #0f172a;
        font-size: 8.5pt;
        margin-bottom: 8px;
    }
    .assigned-table th {
        background: #e2e8f0;
        color: #0f172a;
        border: 1px solid #64748b;
        padding: 3px 4px;
        font-weight: 800;
        text-align: center;
    }
    .assigned-table td {
        border: 1px solid #94a3b8;
        padding: 2px 4px;
        text-align: center;
    }
    .official-footer {
        display: flex;
        justify-content: space-between;
        align-items: center;
        border-top: 1.5px solid #0f172a;
        padding-top: 6px;
        font-size: 9.5pt;
        font-weight: 700;
        color: #334155;
    }
    .print-btn-bar {
        display: none !important;
    }
    .btn-action {
        background: #2563eb;
        color: #ffffff;
        border: none;
        padding: 8px 16px;
        border-radius: 20px;
        font-weight: 800;
        font-size: 11pt;
        cursor: pointer;
        display: flex;
        align-items: center;
        gap: 6px;
    }
    .btn-action:hover {
        background: #1d4ed8;
    }
    
    @media screen {
        .print-btn-bar {
            position: fixed !important;
            bottom: 15px !important;
            left: 50% !important;
            transform: translateX(-50%) !important;
            background: #0f172a !important;
            padding: 8px 18px !important;
            border-radius: 30px !important;
            box-shadow: 0 10px 25px rgba(0,0,0,0.3) !important;
            display: flex !important;
            gap: 12px !important;
            z-index: 9999 !important;
        }
    }

    @media print {
        @page {
            size: A4 landscape;
            margin: 4mm 5mm;
        }
        body {
            background: #ffffff !important;
            padding: 0 !important;
        }
        .print-btn-bar {
            display: none !important;
        }
        .sheet {
            border: 2px solid #0f172a !important;
            box-shadow: none !important;
            margin: 0 0 0 0 !important;
            padding: 4mm !important;
            width: 100% !important;
            max-width: 100% !important;
        }
        * {
            -webkit-print-color-adjust: exact !important;
            print-color-adjust: exact !important;
        }
    }
    """

def find_browser_executable() -> Optional[str]:
    candidates = [
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
        os.path.expandvars(r"%LOCALAPPDATA%\Google\Chrome\Application\chrome.exe"),
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
        os.path.expandvars(r"%LOCALAPPDATA%\Microsoft\Edge\Application\msedge.exe"),
    ]
    for c in candidates:
        if os.path.exists(c):
            return c
    return None

def convert_html_to_pdf(html_file: str, pdf_file: str) -> bool:
    import re
    try:
        with open(html_file, "r", encoding="utf-8") as f:
            html_content = f.read()

        # Strip onscreen interactive button bar so PDF is 100% clean
        clean_html = re.sub(r'<div class="print-btn-bar">[\s\S]*?</div>', '', html_content)
        
        # Save temp clean file for conversion
        clean_temp_file = html_file + ".clean.html"
        with open(clean_temp_file, "w", encoding="utf-8") as f:
            f.write(clean_html)

        try:
            import pymupdf
            doc = pymupdf.open(clean_temp_file)
            pdf_bytes = doc.convert_to_pdf()
            pdf_doc = pymupdf.open("pdf", pdf_bytes)
            pdf_doc.save(pdf_file)
            pdf_doc.close()
            doc.close()
            if os.path.exists(clean_temp_file):
                try:
                    os.remove(clean_temp_file)
                except Exception:
                    pass
            return os.path.exists(pdf_file) and os.path.getsize(pdf_file) > 0
        except Exception as e:
            print(f"PyMuPDF error for {html_file}:", e)
            if os.path.exists(clean_temp_file):
                try:
                    os.remove(clean_temp_file)
                except Exception:
                    pass
            return False
    except Exception as e:
        print(f"PDF conversion general error for {html_file}:", e)
        return False

class ExportEngine:
    def __init__(self):
        self.desktop_dir = os.path.join(os.path.expanduser("~"), "Desktop", "RzzakFet")
        self.classes_dir = os.path.join(self.desktop_dir, "جداول_الاقسام")
        self.teachers_dir = os.path.join(self.desktop_dir, "جداول_الاساتذة")
        self.master_dir = os.path.join(self.desktop_dir, "الجدول_العام")

    def ensure_directories(self):
        os.makedirs(self.desktop_dir, exist_ok=True)
        os.makedirs(self.classes_dir, exist_ok=True)
        os.makedirs(self.teachers_dir, exist_ok=True)
        os.makedirs(self.master_dir, exist_ok=True)

    def _are_cells_mergeable(self, c1: Optional[Dict[str, Any]], c2: Optional[Dict[str, Any]], category: str) -> bool:
        if not c1 or not c2:
            return False
        if c1.get("is_grouped") or c2.get("is_grouped"):
            if c1.get("is_grouped") and c2.get("is_grouped"):
                sg1 = c1.get("subgroups", {})
                sg2 = c2.get("subgroups", {})
                g1_a = sg1.get("G1") or sg1.get("F1", {})
                g1_b = sg2.get("G1") or sg2.get("F1", {})
                g2_a = sg1.get("G2") or sg1.get("F2", {})
                g2_b = sg2.get("G2") or sg2.get("F2", {})
                return (g1_a.get("subject") == g1_b.get("subject") and
                        g1_a.get("room") == g1_b.get("room") and
                        g2_a.get("subject") == g2_b.get("subject") and
                        g2_a.get("room") == g2_b.get("room"))
            return False
        if category == "class":
            return (c1.get("subject") == c2.get("subject") and
                    c1.get("teacher") == c2.get("teacher") and
                    c1.get("room") == c2.get("room"))
        elif category == "teacher":
            return (c1.get("students") == c2.get("students") and
                    c1.get("subject") == c2.get("subject") and
                    c1.get("room") == c2.get("room"))
        else:
            return (c1.get("teacher") == c2.get("teacher") and
                    c1.get("students") == c2.get("students") and
                    c1.get("subject") == c2.get("subject"))

    def render_single_sheet_html(
        self,
        category: str,
        target_name: str,
        grid_data: Dict[str, Dict[str, Any]],
        inst_info: Dict[str, Any],
        break_slots: List[Dict[str, str]],
        generated_date: str,
        wed_settings: Optional[Dict[str, Any]] = None
    ) -> str:
        inst_name = inst_info.get("name", "الثانوية الإعدادية")
        prov_name = inst_info.get("province", "المديرية الإقليمية")
        acad_name = inst_info.get("academy", "الأكاديمية الجهوية")
        acad_year = inst_info.get("academic_year") or inst_info.get("year", "2026-2027")
        gresa_code = inst_info.get("gresa_code") or inst_info.get("gresa", "045892P")
        principal = inst_info.get("principal_name") or inst_info.get("principal", "رئيس المؤسسة")
        logo_src = inst_info.get("logo_base64") or inst_info.get("logo") or "/official_header_logo.png"
        logo_html = f'''
        <div class="official-logo-box" style="text-align:center; margin-bottom:4px;">
            <img src="{logo_src}" alt="شعار الوزارة والمؤسسة" style="max-height:42px; max-width:260px; object-fit:contain; display:inline-block;">
        </div>
        ''' if logo_src else ''

        if category == "class":
            main_title = f"جدول حصص القسم: {target_name}"
        elif category == "teacher":
            main_title = f"جدول حصص الأستاذ(ة): {target_name}"
        else:
            main_title = f"جدول إشغال القاعة: {target_name}"

        total_hours = 0
        assigned_teachers = {}

        rows_html = ""
        for day_obj in DAYS_ORDER:
            day_name = day_obj["name"]
            m_key = day_obj["m"]
            e_key = day_obj["e"]

            m_slots = [(m_key, h) for h in HOURS_4]
            e_slots = [(e_key, h) for h in HOURS_4]
            all_8 = m_slots + e_slots

            cells_html = ""
            k = 0
            while k < len(all_8):
                slot_day, slot_h = all_8[k]
                is_break = any(s.get("day") == slot_day and s.get("hour") == slot_h for s in break_slots)
                cell = (grid_data.get(slot_day) or {}).get(slot_h)

                if category == "teacher" and (slot_day == "الاربعاء م" or slot_day == "الأربعاء م") and wed_settings and wed_settings.get("enabled", True) and not cell:
                    pe_label = wed_settings.get("pe_label", "أنشطة الجمعية الرياضية المدرسية (ASS)")
                    gen_label = wed_settings.get("general_label", "الأنشطة الموازية والدعم التربوي")
                    is_pe = ("بدن" in target_name or "EPS" in target_name or
                             any("بدن" in str(c.get("subject", "")) or "EPS" in str(c.get("subject", ""))
                                 for d in grid_data.values() for c in d.values() if isinstance(c, dict)))
                    if is_pe:
                        cells_html += f'<td class="cell-ass"><div>{pe_label}</div><div style="font-size:7.5pt;opacity:0.85;">رياضة الكم والنخبة</div></td>'
                    else:
                        cells_html += f'<td class="cell-support"><div>{gen_label}</div><div style="font-size:7.5pt;opacity:0.85;">مواكبة وتأطير تربوي</div></td>'
                    k += 1
                    continue

                if category == "class" and (slot_day == "الاربعاء م" or slot_day == "الأربعاء م") and wed_settings and wed_settings.get("enabled", True) and not cell:
                    cells_html += '<td class="cell-break" style="background:#eff6ff;color:#1e40af;font-size:7.5pt;font-weight:bold;text-align:center;"><div>أنشطة الجمعية الرياضية (ASS)</div><div style="font-size:6.5pt;color:#3b82f6;">أنشطة موازية ودعم</div></td>'
                    k += 1
                    continue

                if is_break:
                    cells_html += '<td class="cell-break"></td>'
                    k += 1
                    continue

                if not cell:
                    cells_html += '<td class="cell-empty"></td>'
                    k += 1
                    continue

                # Check if this cell can merge with next slot within the same half-day (0..3 morning, 4..7 afternoon)
                span = 1
                half_day_end = 4 if k < 4 else 8
                if k + 1 < half_day_end:
                    next_day, next_h = all_8[k + 1]
                    next_is_break = any(s.get("day") == next_day and s.get("hour") == next_h for s in break_slots)
                    next_cell = (grid_data.get(next_day) or {}).get(next_h)
                    if not next_is_break and next_cell and self._are_cells_mergeable(cell, next_cell, category):
                        span = 2

                total_hours += span
                subj = cell.get("subject", "")
                teach = cell.get("teacher", "")
                stds = cell.get("students", "")
                room = cell.get("room", "")

                if subj and teach:
                    assigned_teachers[subj] = teach

                style = get_subj_style(subj)
                bg_col = style["print_bg"]
                txt_col = style["print_text"]

                colspan_attr = f' colspan="{span}"' if span > 1 else ''

                if category == "class":
                    if cell.get("is_grouped") and cell.get("subgroups"):
                        g1_inf = cell["subgroups"].get("G1") or cell["subgroups"].get("F1", {})
                        g2_inf = cell["subgroups"].get("G2") or cell["subgroups"].get("F2", {})
                        s1 = g1_inf.get("subject", "فيزياء")
                        r1 = abbreviate_room_name(g1_inf.get("room", "PC"))
                        s2 = g2_inf.get("subject", "علوم")
                        r2 = abbreviate_room_name(g2_inf.get("room", "SVT"))
                        cells_html += f'''
                        <td{colspan_attr} style="padding: 2px !important; background-color: #f8fafc; color: #0f172a; vertical-align: middle;">
                            <div style="background:#e0f2fe; color:#0369a1; border-radius:3px; padding:2px 4px; margin-bottom:2px; font-size:7.5pt; font-weight:800; border: 1px solid #7dd3fc; line-height:1.2; display:flex; justify-content:space-between;">
                                <span>G1: {s1}</span> <span style="font-size:6.5pt; opacity:0.85;">({r1})</span>
                            </div>
                            <div style="background:#fef3c7; color:#b45309; border-radius:3px; padding:2px 4px; font-size:7.5pt; font-weight:800; border: 1px solid #fde68a; line-height:1.2; display:flex; justify-content:space-between;">
                                <span>G2: {s2}</span> <span style="font-size:6.5pt; opacity:0.85;">({r2})</span>
                            </div>
                        </td>
                        '''
                    else:
                        cells_html += f'''
                        <td{colspan_attr} style="background-color: {bg_col}; color: {txt_col};">
                            <div class="subj-cell">{subj}</div>
                            <div class="sub-info">{room}</div>
                        </td>
                        '''
                elif category == "teacher":
                    std_display = stds
                    if "-G1" in stds:
                        std_display = stds.replace("-G1", " (G1)")
                    elif "-G2" in stds:
                        std_display = stds.replace("-G2", " (G2)")
                    elif "-F1" in stds:
                        std_display = stds.replace("-F1", " (G1)")
                    elif "-F2" in stds:
                        std_display = stds.replace("-F2", " (G2)")
                    cls_style = get_class_style(stds)
                    bg_col = cls_style["print_bg"]
                    txt_col = cls_style["print_text"]
                    cells_html += f'''
                    <td{colspan_attr} style="background-color: {bg_col} !important; color: {txt_col} !important;">
                        <div class="subj-cell" style="color: {txt_col} !important; font-weight:900;">{std_display}</div>
                        <div class="sub-info" style="color: {txt_col} !important; opacity:0.95;">{room}</div>
                    </td>
                    '''
                else:
                    cells_html += f'''
                    <td{colspan_attr} style="background-color: {bg_col}; color: {txt_col};">
                        <div class="subj-cell">{teach}</div>
                        <div class="sub-info">{stds}</div>
                    </td>
                    '''

                k += span

            rows_html += f'''
            <tr>
                <td class="day-cell">{day_name}</td>
                {cells_html}
            </tr>
            '''

        bottom_table_html = ""
        if category == "class" and assigned_teachers:
            items = list(assigned_teachers.items())
            pairs_per_row = 4
            assigned_rows = ""
            for i in range(0, len(items), pairs_per_row):
                chunk = items[i:i+pairs_per_row]
                while len(chunk) < pairs_per_row:
                    chunk.append(("-", "-"))
                tds = "".join([f'<td style="background:#f1f5f9;font-weight:bold;">{s}</td><td style="background:#ffffff;">{t}</td>' for s, t in chunk])
                assigned_rows += f'<tr>{tds}</tr>'

            bottom_table_html = f'''
            <div style="margin-top: 4px;">
                <div style="text-align:center; font-weight:800; font-size:8.5pt; margin-bottom:2px; color:#1e293b;">لائحة الأساتذة المسندين حسب المواد</div>
                <table class="assigned-table">
                    <thead>
                        <tr>
                            <th>المادة</th><th>الأستاذ</th>
                            <th>المادة</th><th>الأستاذ</th>
                            <th>المادة</th><th>الأستاذ</th>
                            <th>المادة</th><th>الأستاذ</th>
                        </tr>
                    </thead>
                    <tbody>
                        {assigned_rows}
                    </tbody>
                </table>
            </div>
            '''

        logo_src = inst_info.get("logo_base64") or inst_info.get("logo") or "/official_header_logo.png"
        logo_html = f'''
        <div class="official-logo-box" style="text-align:center; margin-bottom:5px;">
            <img src="{logo_src}" alt="شعار الوزارة والمؤسسة" style="max-height:50px; max-width:300px; object-fit:contain; display:inline-block;">
        </div>
        ''' if logo_src else ''

        return f'''
        <div class="sheet">
            {logo_html}
            <div class="official-header" style="display:flex; justify-content:space-between; align-items:center; border-bottom:1.5px solid #0f172a; padding-bottom:4px; margin-bottom:6px;">
                <div class="header-box" style="text-align:right; font-size:8.5pt; line-height:1.4;">
                    <div style="font-weight:800; color:#1e293b;">{acad_name}</div>
                    <div style="font-weight:700; color:#334155;">{prov_name}</div>
                    <div style="font-weight:900; color:#0f172a; font-size:9.5pt;">{inst_name}</div>
                </div>

                <div class="header-center" style="text-align:center;">
                    <div class="main-title" style="font-size:12pt; font-weight:900; color:#0f172a; margin-bottom:2px;">{main_title}</div>
                    <div style="font-size:8pt; color:#475569;">الموسم الدراسي: <strong>{acad_year}</strong> | رمز مسار: <strong>{gresa_code}</strong></div>
                </div>
            </div>

            <div class="sub-hierarchy">
                <div>{acad_name} - {prov_name} - {inst_name}</div>
                <div>الحصص الأسبوعية: <strong style="background:#0f172a; color:#fff; padding:1px 6px; border-radius:3px; font-family:monospace;">{total_hours} h</strong></div>
            </div>

            <table class="grid-table">
                <thead>
                    <tr>
                        <th style="width:9%;">اليوم</th>
                        {"".join([f'<th>{t}</th>' for t in (inst_info.get("morning_period_times") or ["08:30-09:30", "09:30-10:30", "10:30-11:30", "11:30-12:30"])[:4]])}
                        {"".join([f'<th>{t}</th>' for t in (inst_info.get("afternoon_period_times") or ["14:30-15:30", "15:30-16:30", "16:30-17:30", "17:30-18:30"])[:4]])}
                    </tr>
                </thead>
                <tbody>
                    {rows_html}
                </tbody>
            </table>

            {bottom_table_html}

            <div class="official-footer">
                <div>حرر في: <span style="font-family:monospace;">{generated_date}</span></div>
                <div>توقيع وتأشيرة المعني بالأمر</div>
                <div>خاتم وتأشيرة رئيس المؤسسة</div>
            </div>
        </div>
        '''

    def render_master_table_html(
        self,
        master_data: List[Dict[str, Any]],
        inst_info: Dict[str, Any],
        break_slots: List[Dict[str, str]],
        generated_date: str
    ) -> str:
        inst_name = inst_info.get("name", "الثانوية الإعدادية")
        prov_name = inst_info.get("province", "المديرية الإقليمية")
        acad_name = inst_info.get("academy", "الأكاديمية الجهوية")
        acad_year = inst_info.get("academic_year") or inst_info.get("year", "2026-2027")
        gresa_code = inst_info.get("gresa_code") or inst_info.get("gresa", "045892P")
        principal = inst_info.get("principal_name") or inst_info.get("principal", "رئيس المؤسسة")
        logo_src = inst_info.get("logo_base64") or inst_info.get("logo") or "/official_header_logo.png"
        logo_html = f'''
        <div class="official-logo-box" style="text-align:center; margin-bottom:4px;">
            <img src="{logo_src}" alt="شعار الوزارة والمؤسسة" style="max-height:42px; max-width:260px; object-fit:contain; display:inline-block;">
        </div>
        ''' if logo_src else ''

        rows_html = ""
        DAYS_12 = [
            "الاثنين ص", "الاثنين م",
            "الثلاثاء ص", "الثلاثاء م",
            "الاربعاء ص", "الاربعاء م",
            "الخميس ص", "الخميس م",
            "الجمعة ص", "الجمعة م",
            "السبت ص", "السبت م"
        ]

        for idx, t_row in enumerate(master_data):
            teacher = t_row.get("teacher", "")
            subject = t_row.get("subject", "")
            room = t_row.get("room", "")
            schedule = t_row.get("schedule", {})
            style = get_subj_style(subject)
            bg_subj = style["print_bg"]

            cells = ""
            for day in DAYS_12:
                h_idx = 0
                while h_idx < len(HOURS_4):
                    h = HOURS_4[h_idx]
                    is_break = any(s.get("day") == day and s.get("hour") == h for s in break_slots)
                    cell = (schedule.get(day) or {}).get(h)

                    if is_break and not cell:
                        span = 1
                        while (h_idx + span < len(HOURS_4)):
                            next_h = HOURS_4[h_idx + span]
                            next_break = any(s.get("day") == day and s.get("hour") == next_h for s in break_slots)
                            next_c = (schedule.get(day) or {}).get(next_h)
                            if next_break and not next_c:
                                span += 1
                            else:
                                break
                        colspan_attr = f' colspan="{span}"' if span > 1 else ''
                        cells += f'<td class="break-cell"{colspan_attr} style="background:#cbd5e1; color:#475569; font-size:6pt; font-weight:900;">✕</td>'
                        h_idx += span
                    elif not cell:
                        cells += '<td class="empty-cell" style="background:#f8fafc;"></td>'
                        h_idx += 1
                    else:
                        stds = cell.get("students", "")
                        rm = cell.get("room", "")
                        span = 1
                        while (h_idx + span < len(HOURS_4)):
                            next_h = HOURS_4[h_idx + span]
                            next_is_break = any(s.get("day") == day and s.get("hour") == next_h for s in break_slots)
                            if next_is_break:
                                break
                            next_cell = (schedule.get(day) or {}).get(next_h)
                            if next_cell and next_cell.get("students", "") == stds and next_cell.get("room", "") == rm:
                                span += 1
                            else:
                                break

                        colspan_attr = f' colspan="{span}"' if span > 1 else ''
                        cell_content = format_master_cell_class(stds, rm, span)
                        cells += f'''
                        <td class="m-cell"{colspan_attr} style="background:{bg_subj};" title="{teacher} | {subject} | {stds} ({rm})">
                            {cell_content}
                        </td>
                        '''
                        h_idx += span

            short_home_room = abbreviate_room_name(room)
            rows_html += f'''
            <tr>
                <td style="font-weight:900; background:#f1f5f9; text-align:right; padding:1px 3px; border:0.5px solid #64748b; font-size:6.5pt; white-space:nowrap; overflow:hidden; text-overflow:ellipsis;" title="{teacher}">{teacher}</td>
                <td style="background:#ffffff; text-align:right; padding:1px 2px; border:0.5px solid #64748b; font-size:6pt; white-space:nowrap; overflow:hidden; text-overflow:ellipsis;" title="{subject}">{subject}</td>
                <td style="background:#f8fafc; text-align:center; padding:1px; border:0.5px solid #64748b; font-size:5.5pt; white-space:nowrap; overflow:hidden; text-overflow:ellipsis;" title="{room}">{short_home_room}</td>
                {cells}
            </tr>
            '''

        return f'''<!DOCTYPE html>
<html lang="ar" dir="rtl">
<head>
    <meta charset="UTF-8">
    <title>الجدول العام لاستعمالات الزمن للمدرسين - {inst_name}</title>
    <style id="page-style">
        @page {{
            size: A3 landscape;
            margin: 3mm 4mm;
        }}
    </style>
    <style>
        {get_base_css()}
        
        body {{
            background-color: #0b1120;
            color: #0f172a;
            direction: rtl;
            padding: 10px;
        }}

        .master-sheet {{
            background: #ffffff;
            border: 2px solid #0f172a;
            border-radius: 4px;
            padding: 5px 8px;
            width: 100%;
            margin: 0 auto;
            box-sizing: border-box;
            box-shadow: 0 10px 25px rgba(0,0,0,0.4);
        }}

        .master-table {{
            width: 100%;
            table-layout: fixed;
            border-collapse: collapse;
            border: 1.5px solid #0f172a;
            text-align: center;
            font-size: 6pt;
        }}

        .master-table th {{
            background: #1e293b;
            color: #ffffff;
            border: 0.5px solid #475569;
            padding: 2px 0.5px;
            font-weight: 800;
            overflow: hidden;
            white-space: nowrap;
            text-overflow: ellipsis;
        }}

        .master-table td {{
            border: 0.5px solid #64748b;
            padding: 1px 0.5px;
            height: 25px;
            min-height: 25px;
            vertical-align: middle;
            text-align: center;
            box-sizing: border-box;
        }}

        .master-table td.m-cell {{
            padding: 1px 0.5px !important;
            color: #ffffff !important;
        }}

        .master-table td.break-cell {{
            border: 0.5px solid #94a3b8;
        }}

        .master-table td.empty-cell {{
            border: 0.5px solid #cbd5e1;
        }}

        .c-code {{
            font-size: 5.6pt;
            font-weight: 800;
            line-height: 1.0;
            color: #ffffff;
            display: flex;
            flex-direction: column;
            align-items: center;
            justify-content: center;
            width: 100%;
        }}

        .c-code.full-span {{
            font-size: 6.8pt;
            font-weight: 900;
            white-space: nowrap;
            overflow: hidden;
            text-overflow: ellipsis;
            padding: 0 1px;
            letter-spacing: -0.2px;
        }}

        .c-code.single-span {{
            font-size: 5pt;
            line-height: 0.95;
        }}

        .c-line1 {{
            font-size: 5.1pt;
            font-weight: 800;
            white-space: nowrap;
            letter-spacing: -0.2px;
        }}

        .c-line2 {{
            font-size: 5.3pt;
            font-weight: 900;
            white-space: nowrap;
            background: rgba(0, 0, 0, 0.28);
            border-radius: 2px;
            padding: 0 2px;
            margin-top: 1px;
        }}

        .r-code {{
            font-size: 4.5pt;
            font-weight: 700;
            line-height: 1.0;
            opacity: 0.92;
            color: #ffffff;
            margin-top: 1px;
            white-space: nowrap;
        }}

        @media print {{
            body {{
                background: #ffffff !important;
                padding: 0 !important;
            }}
            .master-sheet {{
                border: 1px solid #0f172a !important;
                padding: 1.5mm !important;
                box-shadow: none !important;
            }}
            .print-btn-bar {{
                display: none !important;
            }}
            tr {{
                page-break-inside: avoid !important;
            }}
            * {{
                -webkit-print-color-adjust: exact !important;
                print-color-adjust: exact !important;
            }}
        }}
    </style>
</head>
<body>
    <div class="print-btn-bar" style="display:flex; flex-wrap:wrap; gap:8px; align-items:center; justify-content:center; padding:8px; background:#0f172a; border-radius:8px; margin-bottom:10px;">
        <button onclick="window.location.href='/export/master_excel'" class="btn-action" style="background:#16a34a; font-weight:800; display:flex; align-items:center; gap:5px;">
            <span>📊 تحميل ملف إكسيل (Excel .xlsx)</span>
        </button>
        <button onclick="window.location.href='/export/master_pdf'" class="btn-action" style="background:#dc2626; font-weight:800; display:flex; align-items:center; gap:5px;">
            <span>📄 تحميل ملف (PDF)</span>
        </button>
        <button onclick="window.location.href='/export/master_html'" class="btn-action" style="background:#0284c7; font-weight:800; display:flex; align-items:center; gap:5px;">
            <span>🌐 تحميل ملف (HTML)</span>
        </button>
        <div style="width:1px; height:24px; background:#334155; margin:0 4px;"></div>
        <button onclick="printA3()" class="btn-action" style="background:#4f46e5; font-weight:800;">🖨️ طباعة الجدول العام (A3 - حجم كامل)</button>
        <button onclick="printA4()" class="btn-action" style="background:#0891b2; font-weight:800;">🖨️ طباعة الجدول العام (A4 - مكيف تلقائياً)</button>
        <button onclick="window.close()" class="btn-action" style="background:#475569;">✕ إغلاق</button>
    </div>

    <div class="master-sheet">
        {logo_html}
        <div class="official-header" style="display:flex; justify-content:space-between; align-items:center; padding-bottom:3px; margin-bottom:4px; border-bottom:1.5px solid #0f172a;">
            <div class="header-box" style="text-align:right; font-size:7.5pt; line-height:1.35;">
                <div style="font-weight:800; color:#1e293b;">{acad_name} | {prov_name}</div>
                <div style="font-weight:900; font-size:8.5pt; color:#0f172a;">{inst_name}</div>
            </div>

            <div class="header-center" style="text-align:center;">
                <div class="main-title" style="font-size:10pt; font-weight:900; padding:2px 12px; margin-bottom:1px;">📋 الجدول العام لاستعمالات الزمن للمدرسين</div>
                <div style="font-size:7.5pt; color:#475569;">الموسم الدراسي: <strong>{acad_year}</strong> | رمز المؤسسة: <strong>{gresa_code}</strong> | رئيس المؤسسة: <strong>{principal}</strong></div>
            </div>
        </div>

        <table class="master-table">
            <thead>
                <tr>
                    <th rowspan="2" style="width:7.5%;">الأستاذ(ة)</th>
                    <th rowspan="2" style="width:5.5%;">المادة</th>
                    <th rowspan="2" style="width:3.5%;">القاعة</th>
                    <th colspan="4" style="background:#0c4a6e; width:6.9%;">الاثنين ص</th>
                    <th colspan="4" style="width:6.9%;">الاثنين م</th>
                    <th colspan="4" style="background:#0c4a6e; width:6.9%;">الثلاثاء ص</th>
                    <th colspan="4" style="width:6.9%;">الثلاثاء م</th>
                    <th colspan="4" style="background:#0c4a6e; width:6.9%;">الاربعاء ص</th>
                    <th colspan="4" style="width:6.9%;">الاربعاء م</th>
                    <th colspan="4" style="background:#0c4a6e; width:6.9%;">الخميس ص</th>
                    <th colspan="4" style="width:6.9%;">الخميس م</th>
                    <th colspan="4" style="background:#0c4a6e; width:6.9%;">الجمعة ص</th>
                    <th colspan="4" style="width:6.9%;">الجمعة م</th>
                    <th colspan="4" style="background:#0c4a6e; width:6.9%;">السبت ص</th>
                    <th colspan="4" style="width:6.9%;">السبت م</th>
                </tr>
                <tr style="font-size:5pt; font-family:monospace;">
                    <th style="width:1.72%;">1</th><th style="width:1.72%;">2</th><th style="width:1.72%;">3</th><th style="width:1.72%;">4</th>
                    <th style="width:1.72%;">1</th><th style="width:1.72%;">2</th><th style="width:1.72%;">3</th><th style="width:1.72%;">4</th>
                    <th style="width:1.72%;">1</th><th style="width:1.72%;">2</th><th style="width:1.72%;">3</th><th style="width:1.72%;">4</th>
                    <th style="width:1.72%;">1</th><th style="width:1.72%;">2</th><th style="width:1.72%;">3</th><th style="width:1.72%;">4</th>
                    <th style="width:1.72%;">1</th><th style="width:1.72%;">2</th><th style="width:1.72%;">3</th><th style="width:1.72%;">4</th>
                    <th style="width:1.72%;">1</th><th style="width:1.72%;">2</th><th style="width:1.72%;">3</th><th style="width:1.72%;">4</th>
                    <th style="width:1.72%;">1</th><th style="width:1.72%;">2</th><th style="width:1.72%;">3</th><th style="width:1.72%;">4</th>
                    <th style="width:1.72%;">1</th><th style="width:1.72%;">2</th><th style="width:1.72%;">3</th><th style="width:1.72%;">4</th>
                    <th style="width:1.72%;">1</th><th style="width:1.72%;">2</th><th style="width:1.72%;">3</th><th style="width:1.72%;">4</th>
                    <th style="width:1.72%;">1</th><th style="width:1.72%;">2</th><th style="width:1.72%;">3</th><th style="width:1.72%;">4</th>
                    <th style="width:1.72%;">1</th><th style="width:1.72%;">2</th><th style="width:1.72%;">3</th><th style="width:1.72%;">4</th>
                    <th style="width:1.72%;">1</th><th style="width:1.72%;">2</th><th style="width:1.72%;">3</th><th style="width:1.72%;">4</th>
                </tr>
            </thead>
            <tbody>
                {rows_html}
            </tbody>
        </table>

        <div class="official-footer" style="margin-top:6px;">
            <div>حرر في: <span style="font-family:monospace;">{generated_date}</span></div>
            <div>مجموع المدرسين: {len(master_data)} أستاذ(ة)</div>
            <div>خاتم وتأشيرة رئيس المؤسسة</div>
        </div>
    </div>

    <script>
    function printA3() {{
        document.getElementById('page-style').innerHTML = '@page {{ size: A3 landscape; margin: 3mm 4mm; }}';
        window.print();
    }}
    function printA4() {{
        document.getElementById('page-style').innerHTML = `
            @page {{ size: A4 landscape; margin: 1.5mm 2mm; }}
            .master-sheet {{ padding: 1mm !important; border: 0.5px solid #0f172a !important; }}
            .official-header {{ padding-bottom: 1px !important; margin-bottom: 2px !important; }}
            .header-box {{ font-size: 6pt !important; line-height: 1.1 !important; }}
            .main-title {{ font-size: 7.5pt !important; padding: 1px 6px !important; margin-bottom: 1px !important; }}
            .master-table th {{ padding: 0.5px !important; font-size: 4.2pt !important; }}
            .master-table td {{ height: 14px !important; max-height: 14px !important; font-size: 4.2pt !important; }}
            .c-code {{ font-size: 4.8pt !important; letter-spacing: -0.4px !important; line-height: 0.95 !important; }}
            .r-code {{ font-size: 3.5pt !important; line-height: 0.95 !important; }}
            .official-footer {{ font-size: 6pt !important; margin-top: 2px !important; padding-top: 2px !important; }}
        `;
        window.print();
    }}
    </script>
</body>
</html>
'''

    def generate_master_table_excel(
        self,
        master_data: List[Dict[str, Any]],
        inst_info: Dict[str, Any],
        break_slots: List[Dict[str, str]],
        output_path: Optional[str] = None
    ) -> str:
        if not output_path:
            self.ensure_directories()
            output_path = os.path.join(self.master_dir, "الجدول_العام_لاستعمالات_الزمن_للمدرسين.xlsx")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "الجدول العام للمدرسين"
        ws.views.sheetView[0].rightToLeft = True

        inst_name = inst_info.get("name", "الثانوية الإعدادية")
        prov_name = inst_info.get("province", "المديرية الإقليمية")
        acad_name = inst_info.get("academy", "الأكاديمية الجهوية")
        acad_year = inst_info.get("academic_year") or inst_info.get("year", "2026-2027")
        gresa_code = inst_info.get("gresa_code") or inst_info.get("gresa", "045892P")
        principal = inst_info.get("principal_name") or inst_info.get("principal", "رئيس المؤسسة")
        logo_src = inst_info.get("logo_base64") or inst_info.get("logo") or "/official_header_logo.png"
        logo_html = f'''
        <div class="official-logo-box" style="text-align:center; margin-bottom:4px;">
            <img src="{logo_src}" alt="شعار الوزارة والمؤسسة" style="max-height:42px; max-width:260px; object-fit:contain; display:inline-block;">
        </div>
        ''' if logo_src else ''

        thin_side = Side(style='thin', color='64748B')
        cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        # 1. Official Header
        ws.cell(row=1, column=1, value="المملكة المغربية").font = Font(name="Calibri", size=11, bold=True, color="0F172A")
        ws.cell(row=1, column=1).alignment = Alignment(horizontal="right", vertical="center")
        
        ws.merge_cells(start_row=1, start_column=10, end_row=1, end_column=42)
        c_title = ws.cell(row=1, column=10, value="📋 الجدول العام لاستعمالات الزمن للمدرسين")
        c_title.font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
        c_title.fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        c_title.alignment = Alignment(horizontal="center", vertical="center")

        ws.cell(row=1, column=45, value=f"الموسم الدراسي: {acad_year}").font = Font(name="Calibri", size=10, bold=True)
        ws.cell(row=1, column=45).alignment = Alignment(horizontal="left", vertical="center")

        ws.cell(row=2, column=1, value="وزارة التربية الوطنية والتعليم الأولي والرياضة").font = Font(name="Calibri", size=9, bold=True, color="334155")
        ws.cell(row=2, column=1).alignment = Alignment(horizontal="right", vertical="center")

        ws.merge_cells(start_row=2, start_column=10, end_row=2, end_column=42)
        c_inst = ws.cell(row=2, column=10, value=inst_name)
        c_inst.font = Font(name="Calibri", size=11, bold=True, color="1E293B")
        c_inst.alignment = Alignment(horizontal="center", vertical="center")

        ws.cell(row=2, column=45, value=f"رمز المؤسسة: {gresa_code}").font = Font(name="Calibri", size=9.5, bold=True, color="475569")
        ws.cell(row=2, column=45).alignment = Alignment(horizontal="left", vertical="center")

        ws.cell(row=3, column=1, value=f"{acad_name} | {prov_name}").font = Font(name="Calibri", size=8.5, color="475569")
        ws.cell(row=3, column=1).alignment = Alignment(horizontal="right", vertical="center")

        ws.merge_cells(start_row=3, start_column=10, end_row=3, end_column=42)
        c_stats = ws.cell(row=3, column=10, value=f"مجموع المدرسين: {len(master_data)} أستاذ(ة)")
        c_stats.font = Font(name="Calibri", size=9.5, bold=True, color="334155")
        c_stats.alignment = Alignment(horizontal="center", vertical="center")

        ws.cell(row=3, column=45, value=f"رئيس المؤسسة: {principal}").font = Font(name="Calibri", size=9, bold=True, color="475569")
        ws.cell(row=3, column=45).alignment = Alignment(horizontal="left", vertical="center")

        # 2. Table Header
        for col_idx, col_name in [(1, "الأستاذ(ة)"), (2, "المادة"), (3, "القاعة")]:
            c = ws.cell(row=5, column=col_idx, value=col_name)
            c.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
            c.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = cell_border
            ws.merge_cells(start_row=5, start_column=col_idx, end_row=6, end_column=col_idx)

        DAYS_12 = [
            "الاثنين ص", "الاثنين م",
            "الثلاثاء ص", "الثلاثاء م",
            "الاربعاء ص", "الاربعاء م",
            "الخميس ص", "الخميس م",
            "الجمعة ص", "الجمعة م",
            "السبت ص", "السبت م"
        ]

        col_pointer = 4
        for d_idx, day in enumerate(DAYS_12):
            is_morning = "ص" in day
            day_bg = "0C4A6E" if is_morning else "1E293B"
            c_day = ws.cell(row=5, column=col_pointer, value=day)
            c_day.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
            c_day.fill = PatternFill(start_color=day_bg, end_color=day_bg, fill_type="solid")
            c_day.alignment = Alignment(horizontal="center", vertical="center")
            c_day.border = cell_border
            ws.merge_cells(start_row=5, start_column=col_pointer, end_row=5, end_column=col_pointer + 3)

            for p_num in range(1, 5):
                col_curr = col_pointer + p_num - 1
                c_p = ws.cell(row=6, column=col_curr, value=p_num)
                c_p.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
                c_p.fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
                c_p.alignment = Alignment(horizontal="center", vertical="center")
                c_p.border = cell_border

            col_pointer += 4

        # 3. Data Rows
        current_row = 7
        for t_idx, t_row in enumerate(master_data):
            teacher = t_row.get("teacher", "")
            subject = t_row.get("subject", "")
            room = t_row.get("room", "")
            schedule = t_row.get("schedule", {})
            style = get_subj_style(subject)
            bg_hex = style["print_bg"].lstrip("#").upper()

            c_t = ws.cell(row=current_row, column=1, value=teacher)
            c_t.font = Font(name="Calibri", size=9.5, bold=True, color="0F172A")
            c_t.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
            c_t.alignment = Alignment(horizontal="right", vertical="center")
            c_t.border = cell_border

            c_s = ws.cell(row=current_row, column=2, value=subject)
            c_s.font = Font(name="Calibri", size=9, bold=True, color="1E293B")
            c_s.fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
            c_s.alignment = Alignment(horizontal="right", vertical="center")
            c_s.border = cell_border

            c_r = ws.cell(row=current_row, column=3, value=abbreviate_room_name(room))
            c_r.font = Font(name="Calibri", size=8.5, bold=True, color="334155")
            c_r.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
            c_r.alignment = Alignment(horizontal="center", vertical="center")
            c_r.border = cell_border

            col_curr = 4
            for day in DAYS_12:
                h_idx = 0
                while h_idx < len(HOURS_4):
                    h = HOURS_4[h_idx]
                    is_break = any(s.get("day") == day and s.get("hour") == h for s in break_slots)
                    cell = (schedule.get(day) or {}).get(h)

                    if is_break and not cell:
                        span = 1
                        while (h_idx + span < len(HOURS_4)):
                            next_h = HOURS_4[h_idx + span]
                            next_break = any(s.get("day") == day and s.get("hour") == next_h for s in break_slots)
                            next_c = (schedule.get(day) or {}).get(next_h)
                            if next_break and not next_c:
                                span += 1
                            else:
                                break
                        c_cell = ws.cell(row=current_row, column=col_curr, value="✕")
                        c_cell.font = Font(name="Calibri", size=9, bold=True, color="475569")
                        c_cell.fill = PatternFill(start_color="CBD5E1", end_color="CBD5E1", fill_type="solid")
                        c_cell.alignment = Alignment(horizontal="center", vertical="center")
                        c_cell.border = cell_border
                        if span > 1:
                            ws.merge_cells(start_row=current_row, start_column=col_curr, end_row=current_row, end_column=col_curr + span - 1)
                            for extra_col in range(col_curr, col_curr + span):
                                ws.cell(row=current_row, column=extra_col).border = cell_border
                        col_curr += span
                        h_idx += span
                    elif not cell:
                        c_cell = ws.cell(row=current_row, column=col_curr, value="")
                        c_cell.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
                        c_cell.border = cell_border
                        col_curr += 1
                        h_idx += 1
                    else:
                        stds = cell.get("students", "")
                        rm = cell.get("room", "")
                        span = 1
                        while (h_idx + span < len(HOURS_4)):
                            next_h = HOURS_4[h_idx + span]
                            next_is_break = any(s.get("day") == day and s.get("hour") == next_h for s in break_slots)
                            if next_is_break:
                                break
                            next_cell = (schedule.get(day) or {}).get(next_h)
                            if next_cell and next_cell.get("students", "") == stds and next_cell.get("room", "") == rm:
                                span += 1
                            else:
                                break
                        
                        cell_text = f"{stds}\n({abbreviate_room_name(rm)})" if rm else stds
                        c_cell = ws.cell(row=current_row, column=col_curr, value=cell_text)
                        c_cell.font = Font(name="Calibri", size=8.5 if span == 1 else 9.5, bold=True, color="FFFFFF")
                        c_cell.fill = PatternFill(start_color=bg_hex, end_color=bg_hex, fill_type="solid")
                        c_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                        c_cell.border = cell_border
                        if span > 1:
                            ws.merge_cells(start_row=current_row, start_column=col_curr, end_row=current_row, end_column=col_curr + span - 1)
                            for extra_col in range(col_curr, col_curr + span):
                                ws.cell(row=current_row, column=extra_col).border = cell_border
                        col_curr += span
                        h_idx += span

            ws.row_dimensions[current_row].height = 28
            current_row += 1

        # Column widths
        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 14
        ws.column_dimensions["C"].width = 9
        for c_idx in range(4, 52):
            col_letter = get_column_letter(c_idx)
            ws.column_dimensions[col_letter].width = 6.2

        ws.row_dimensions[1].height = 24
        ws.row_dimensions[2].height = 20
        ws.row_dimensions[3].height = 18
        ws.row_dimensions[4].height = 8
        ws.row_dimensions[5].height = 22
        ws.row_dimensions[6].height = 18

        # Footer
        ws.cell(row=current_row + 1, column=1, value=f"حرر في: {time.strftime('%Y-%m-%d')}").font = Font(name="Calibri", size=9, bold=True)
        ws.cell(row=current_row + 1, column=45, value="خاتم وتأشيرة رئيس المؤسسة").font = Font(name="Calibri", size=9, bold=True)

        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

        wb.save(output_path)
        return output_path

    def export_all(self, timetable_data: Dict[str, Any], app_state_data: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
        self.ensure_directories()
        
        inst_info = timetable_data.get("institution_info", {})
        if app_state_data and "institution" in app_state_data:
            inst_info.update(app_state_data["institution"])
        
        raw_breaks = (app_state_data.get("institution", {}).get("break_time_slots") if app_state_data else None) or [
            {"day": "الاربعاء م", "hour": "ح 1"}, {"day": "الاربعاء م", "hour": "ح 2"}, {"day": "الاربعاء م", "hour": "ح 3"}, {"day": "الاربعاء م", "hour": "ح 4"},
            {"day": "السبت م", "hour": "ح 1"}, {"day": "السبت م", "hour": "ح 2"}, {"day": "السبت م", "hour": "ح 3"}, {"day": "السبت م", "hour": "ح 4"}
        ]
        break_slots = [{"day": s.get("day", "").replace("الأربعاء", "الاربعاء").replace("الإثنين", "الاثنين"), "hour": s.get("hour", "")} for s in raw_breaks]
        
        wed_settings = app_state_data.get("wednesday_settings", {}) if app_state_data else {}
        gen_date = timetable_data.get("generated_at", time.strftime("%Y-%m-%d"))

        classes_dict = timetable_data.get("student_timetables", {})
        teachers_dict = timetable_data.get("teacher_timetables", {})
        master_data = timetable_data.get("master_teachers_data", [])

        exported_classes = []
        all_classes_sheets = []

        conversion_tasks = []

        # 1. Classes Export (HTML + PDF)
        for class_name, grid in classes_dict.items():
            sheet_html = self.render_single_sheet_html(
                "class", class_name, grid, inst_info, break_slots, gen_date, wed_settings
            )
            all_classes_sheets.append(sheet_html)

            single_html = f'''<!DOCTYPE html>
<html lang="ar" dir="rtl">
<head>
    <meta charset="UTF-8">
    <title>جدول حصص القسم {class_name}</title>
    <style>{get_base_css()}</style>
</head>
<body>
    <div class="print-btn-bar">
        <button onclick="window.print()" class="btn-action">🖨️ طباعة جدول القسم (A4)</button>
        <button onclick="window.close()" class="btn-action" style="background:#475569;">✕ إغلاق</button>
    </div>
    {sheet_html}
</body>
</html>'''
            safe_name = "".join([c for c in class_name if c.isalnum() or c in (' ', '_', '-')]).strip().replace(' ', '_')
            html_file = os.path.join(self.classes_dir, f"جدول_القسم_{safe_name}.html")
            pdf_file = os.path.join(self.classes_dir, f"جدول_القسم_{safe_name}.pdf")
            with open(html_file, "w", encoding="utf-8") as f:
                f.write(single_html)
            conversion_tasks.append((html_file, pdf_file))
            exported_classes.append(pdf_file)

        # Collective classes HTML + PDF
        all_classes_html = f'''<!DOCTYPE html>
<html lang="ar" dir="rtl">
<head>
    <meta charset="UTF-8">
    <title>جميع جداول الأقسام ({len(classes_dict)} قسماً)</title>
    <style>{get_base_css()}</style>
</head>
<body>
    <div class="print-btn-bar">
        <button onclick="window.print()" class="btn-action">🖨️ طباعة جميع جداول الأقسام ({len(classes_dict)} قسماً)</button>
        <button onclick="window.close()" class="btn-action" style="background:#475569;">✕ إغلاق</button>
    </div>
    {" ".join(all_classes_sheets)}
</body>
</html>'''
        all_classes_path = os.path.join(self.classes_dir, "00_طباعة_جميع_جداول_الأقسام.html")
        all_classes_pdf = os.path.join(self.classes_dir, "جميع_جداول_الأقسام.pdf")
        with open(all_classes_path, "w", encoding="utf-8") as f:
            f.write(all_classes_html)
        conversion_tasks.append((all_classes_path, all_classes_pdf))

        # 2. Teachers Export (HTML + PDF)
        exported_teachers = []
        all_teachers_sheets = []
        for teacher_name, grid in teachers_dict.items():
            sheet_html = self.render_single_sheet_html(
                "teacher", teacher_name, grid, inst_info, break_slots, gen_date, wed_settings
            )
            all_teachers_sheets.append(sheet_html)

            single_html = f'''<!DOCTYPE html>
<html lang="ar" dir="rtl">
<head>
    <meta charset="UTF-8">
    <title>جدول حصص الأستاذ {teacher_name}</title>
    <style>{get_base_css()}</style>
</head>
<body>
    <div class="print-btn-bar">
        <button onclick="window.print()" class="btn-action">🖨️ طباعة جدول الأستاذ (A4)</button>
        <button onclick="window.close()" class="btn-action" style="background:#475569;">✕ إغلاق</button>
    </div>
    {sheet_html}
</body>
</html>'''
            safe_name = "".join([c for c in teacher_name if c.isalnum() or c in (' ', '_', '-')]).strip().replace(' ', '_')
            html_file = os.path.join(self.teachers_dir, f"جدول_الأستاذ_{safe_name}.html")
            pdf_file = os.path.join(self.teachers_dir, f"جدول_الأستاذ_{safe_name}.pdf")
            with open(html_file, "w", encoding="utf-8") as f:
                f.write(single_html)
            conversion_tasks.append((html_file, pdf_file))
            exported_teachers.append(pdf_file)

        # Collective teachers HTML + PDF
        all_teachers_html = f'''<!DOCTYPE html>
<html lang="ar" dir="rtl">
<head>
    <meta charset="UTF-8">
    <title>جميع جداول الأساتذة ({len(teachers_dict)} أستاذاً)</title>
    <style>{get_base_css()}</style>
</head>
<body>
    <div class="print-btn-bar">
        <button onclick="window.print()" class="btn-action">🖨️ طباعة جميع جداول الأساتذة ({len(teachers_dict)} أستاذاً)</button>
        <button onclick="window.close()" class="btn-action" style="background:#475569;">✕ إغلاق</button>
    </div>
    {" ".join(all_teachers_sheets)}
</body>
</html>'''
        all_teachers_path = os.path.join(self.teachers_dir, "00_طباعة_جميع_جداول_الأساتذة.html")
        all_teachers_pdf = os.path.join(self.teachers_dir, "جميع_جداول_الأساتذة.pdf")
        with open(all_teachers_path, "w", encoding="utf-8") as f:
            f.write(all_teachers_html)
        conversion_tasks.append((all_teachers_path, all_teachers_pdf))

        # 3. Master Table HTML + PDF + Excel (.xlsx)
        master_html = self.render_master_table_html(master_data, inst_info, break_slots, gen_date)
        master_path = os.path.join(self.master_dir, "الجدول_العام_لاستعمالات_الزمن_للمدرسين.html")
        master_pdf = os.path.join(self.master_dir, "الجدول_العام_لاستعمالات_الزمن_للمدرسين.pdf")
        master_xlsx = os.path.join(self.master_dir, "الجدول_العام_لاستعمالات_الزمن_للمدرسين.xlsx")

        with open(master_path, "w", encoding="utf-8") as f:
            f.write(master_html)
        conversion_tasks.append((master_path, master_pdf))

        try:
            self.generate_master_table_excel(master_data, inst_info, break_slots, master_xlsx)
        except Exception as e:
            print("Master table Excel generation note:", e)

        # Concurrent PDF generation for maximum speed
        try:
            with ThreadPoolExecutor(max_workers=6) as executor:
                futures = [executor.submit(convert_html_to_pdf, h, p) for h, p in conversion_tasks]
                for fut in futures:
                    fut.result()
        except Exception as e:
            print("Concurrent PDF conversion note:", e)

        return {
            "success": True,
            "desktop_path": self.desktop_dir,
            "classes_folder": self.classes_dir,
            "teachers_folder": self.teachers_dir,
            "master_folder": self.master_dir,
            "total_classes": len(classes_dict),
            "total_teachers": len(teachers_dict),
            "all_classes_file": all_classes_path,
            "all_classes_pdf": all_classes_pdf,
            "all_teachers_file": all_teachers_path,
            "all_teachers_pdf": all_teachers_pdf,
            "master_file": master_path,
            "master_pdf": master_pdf,
            "master_xlsx": master_xlsx
        }

    def open_desktop_folder(self):
        self.ensure_directories()
        try:
            folder_path = os.path.normpath(self.desktop_dir)
            if os.name == 'nt':
                subprocess.Popen(['explorer.exe', folder_path])
            else:
                subprocess.Popen(['xdg-open', folder_path])
            return {"success": True, "path": folder_path}
        except Exception as e:
            return {"success": False, "error": str(e), "path": self.desktop_dir}
